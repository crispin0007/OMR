"""A test series: several exams taken by the same students, with a rank on their combined score.

Each exam is graded as on its own Results tab. A student who has no graded sheet in an exam scores nothing
for it, so the series rank rewards sitting every test; the number of tests taken is shown beside it. A series
can weigh its exams differently (a final mock counting double) and count only each student's best N tests.
"""

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .branding import Institute
from .export import grade, questions
from .grading import DROPPED, full_marks, marks_for, rank
from .store import Store


@dataclass
class SeriesRow:
    roll_number: str
    name: str
    class_name: str
    scores: list[float | None]  # per exam, oldest first; None when the student has no graded sheet in it
    percents: list[float | None]
    total: float = 0.0  # the weighted scores of the tests counted
    average: float = 0.0  # weighted mean percentage over the tests counted
    taken: int = 0
    rank: int = 0
    batch: str = ""
    exam_ranks: list[int | None] = field(default_factory=list)  # rank in each exam on its own
    counted: list[bool] = field(default_factory=list)  # which tests count: all taken, or the best N
    movement: int | None = None  # places gained since the test before the latest (negative: lost); None at first
    improvement: float | None = None  # latest percentage taken minus the first; None with fewer than 2 tests
    sections: dict[str, float] = field(default_factory=dict)  # percentage per section name over every test taken
    _section_marks: dict[str, list[float]] = field(default_factory=dict, repr=False)


def series_results(store: Store, series: str) -> tuple[list[str], list[SeriesRow], list[str]]:
    """Exam names oldest first, one row per student best first, and the exams that couldn't be graded and why."""
    exams = [store.exam(i) for i in store.series_exams(series)]
    settings = store.series_settings(series)
    weights = [settings["weights"].get(str(e.id), 1.0) for e in exams]
    rows: dict[str, SeriesRow] = {}  # by student (`Store.person`), so a new roll number in each test is one row
    person = store.person()
    problems = []
    blank = lambda: SeriesRow("", "", "", [None] * len(exams), [None] * len(exams),  # noqa: E731
                              exam_ranks=[None] * len(exams))
    for n, e in enumerate(exams):
        students = store.students(e.id)
        try:
            graded = grade(e.template, store.pages(e.id), e.answer_key, e.correct, e.wrong, e.blank, students,
                           e.sections, doubled_wrong=e.doubled_wrong)
        except ValueError as err:
            problems.append(f"{e.name}: {err}")
            continue
        if not graded:
            problems.append(f"{e.name}: no results yet (no answer key, or no sheets read)")
        qs = questions(e.template)
        marks = marks_for(qs, e.correct, e.wrong, e.blank, e.sections, doubled_wrong=e.doubled_wrong)
        sat = []
        for g in graded:  # best first, so a roll number read twice keeps its better sheet
            roll = g.page.result.roll_number
            if "?" in roll:
                continue  # an unread roll number belongs to nobody yet; Review fixes it
            row = rows.setdefault(person(e.id, roll), blank())
            row.roll_number = roll  # the latest test's
            if row.scores[n] is not None:
                continue
            out_of = full_marks(g.key, marks)
            row.scores[n] = g.score
            row.percents[n] = round(100 * g.score / out_of, 1) if out_of else None
            sat.append(row)
            for sec in e.sections:
                most = sum(marks[i][0] for i, q in enumerate(qs) if sec.first <= q <= sec.last and g.key[i] != DROPPED)
                got = row._section_marks.setdefault(sec.name, [0.0, 0.0])
                got[0] += g.sections.get(sec.name, 0.0)
                got[1] += most
        for row, (r, _) in zip(sat, rank([row.scores[n] for row in sat])):
            row.exam_ranks[n] = r
        for s in students.values():  # names come from the latest list; listed students without a sheet count too
            row = rows.setdefault(person(e.id, s.roll_number), blank())
            row.roll_number, row.name, row.class_name, row.batch = s.roll_number, s.name, s.class_name, s.batch
    out = list(rows.values())
    for row in out:
        taken = [i for i, s in enumerate(row.scores) if s is not None]
        best = sorted(taken, key=lambda i: -(row.percents[i] or 0))
        keep = set(best[:settings["best_of"]] if settings["best_of"] else taken)
        row.counted = [i in keep for i in range(len(exams))]
        row.taken = len(taken)
        row.total = round(sum(weights[i] * row.scores[i] for i in keep), 2)
        weight = sum(weights[i] for i in keep if row.percents[i] is not None)
        row.average = round(sum(weights[i] * row.percents[i] for i in keep if row.percents[i] is not None)
                            / weight, 1) if weight else 0.0
        if len(taken) >= 2 and row.percents[taken[0]] is not None and row.percents[taken[-1]] is not None:
            row.improvement = round(row.percents[taken[-1]] - row.percents[taken[0]], 1)
        row.sections = {name: round(100 * got / most, 1) for name, (got, most) in row._section_marks.items() if most}
    if len(exams) >= 2:  # places gained or lost with the latest test
        before = _ranks(out, weights[:-1], settings["best_of"])
        now = _ranks(out, weights, settings["best_of"])
        for row in out:
            if row.scores[-1] is not None or row.scores[-2] is not None:
                row.movement = before[id(row)] - now[id(row)]
    out.sort(key=lambda r: (-r.total, r.roll_number))
    for row, (r, _) in zip(out, rank([r.total for r in out])):
        row.rank = r
    return [e.name for e in exams], out, problems


def _ranks(rows: list[SeriesRow], weights: list[float], best_of: int) -> dict[int, int]:
    """Each student's series rank (by the row's id()) counting only the first len(weights) tests."""
    totals = []
    for row in rows:
        taken = [i for i in range(len(weights)) if row.scores[i] is not None]
        keep = sorted(taken, key=lambda i: -(row.percents[i] or 0))[:best_of] if best_of else taken
        totals.append(sum(weights[i] * row.scores[i] for i in keep))
    return {id(row): r for row, (r, _) in zip(rows, rank(totals))}


def exam_summary(exams: list[str], rows: list[SeriesRow]) -> list[tuple[str, int, float, float]]:
    """(exam, students who sat it, average %, top %) per exam, for how the class did across the series."""
    out = []
    for n, name in enumerate(exams):
        p = [r.percents[n] for r in rows if r.percents[n] is not None]
        out.append((name, len(p), round(sum(p) / len(p), 1) if p else 0.0, max(p) if p else 0.0))
    return out


def write_series_excel(path: str | Path, series: str, exams: list[str], rows: list[SeriesRow],
                       institute: Institute | None = None):
    """The series merit list (rank, student, each test's score, total, tests taken, average percentage and the
    places gained with the latest test), and how the class did in each test and section."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Series merit list"
    ws.append([f"{institute.name}: {series}" if institute and institute.name else series])
    ws["A1"].font = Font(bold=True, size=14, color=institute.colour.lstrip("#") if institute else None)
    if institute:
        wb.properties.creator = institute.name
    ws.append(["Rank", "Roll number", "Name", "Class", *exams, "Total", "Tests taken", "Average %", "Movement",
               "Improvement %"])
    for r in rows:
        ws.append([r.rank, r.roll_number, r.name, r.class_name, *("absent" if s is None else s for s in r.scores),
                   r.total, r.taken, r.average, r.movement, r.improvement])
    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A3"
    analysis = wb.create_sheet("Analysis")
    analysis.append(["Test", "Students", "Average %", "Top %"])
    for line in exam_summary(exams, rows):
        analysis.append(list(line))
    names = list(dict.fromkeys(name for r in rows for name in r.sections))
    if names:
        analysis.append([])
        analysis.append(["Section", "Class average %"])
        for name in names:
            p = [r.sections[name] for r in rows if name in r.sections]
            analysis.append([name, round(sum(p) / len(p), 1)])
    for sheet in (ws, analysis):
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column[1:]) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 6), 40)
    wb.save(path)
