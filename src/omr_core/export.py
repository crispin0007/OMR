"""Results for an exam: grade every page, and export a merit list (Excel) and scorecards (PDF) (US3).

The Excel file also has an item analysis: how each question went across the class. When questions have
topics, each student's share right per topic goes on their scorecard and on a Topics sheet.
"""

import csv
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas

from .grading import BONUS, DROPPED, Section, changed_questions, marks_for, parse_key, rank, score, status
from .store import StoredPage, to_review
from .students import Student
from .template import Template


@dataclass
class Graded:
    page: StoredPage
    key: list[str]
    score: float
    right: int
    wrong: int
    blank: int
    rank: int = 0
    percentile: float = 0.0
    to_check: bool = False
    student: Student | None = None  # None when the exam has no student list or the roll number isn't on it
    sections: dict[str, float] = field(default_factory=dict)  # score per section, in the exam's section order
    omr: float = 0.0  # the score from the sheet; `score` adds the written marks
    written: float | None = None  # None when the exam has no written part or this student's aren't typed yet
    topics: dict[str, tuple[int, int]] = field(default_factory=dict)  # (right, counted) per topic, in question order


def questions(t: Template) -> list[int]:
    return t.questions()


def single_questions(t: Template) -> set[int]:
    """Questions printed for one answer only, where two filled bubbles are invalid."""
    return {g.first + r for g in t.grids if g.kind == "questions" and not g.multi for r in range(g.rows)}


def grade(t: Template, pages: list[StoredPage], answer_key: str, correct: float, wrong: float,
          blank: float, students: dict[str, Student] | None = None, sections: list[Section] = (),
          written: dict[str, float] | None = None, topics: dict[int, str] | None = None) -> list[Graded]:
    """Every page that can be scored, best first, with rank and percentile.

    `written` holds typed marks for a written part by roll number; they add to the sheet's score.

    Pages that couldn't be read, or whose set has no key, are left out; they show in review instead.
    Raises ValueError if the key doesn't fit the sheet.
    """
    qs = questions(t)
    keys = parse_key(answer_key, len(qs), t.answer_options())
    marks = marks_for(qs, correct, wrong, blank, sections, single_questions(t))
    students = students or {}
    flagged = {i for i, _, _ in to_review(pages, students)}
    out = []
    for i, p in enumerate(pages):
        key = p.result and keys.get(p.result.set_code, keys.get(""))
        if not key:
            continue
        answers = [p.result.answers[q] for q in qs]
        counts = [status(a, k) for a, k in zip(answers, key)]
        per_section = {}
        for sec in sections:
            part = [n for n, q in enumerate(qs) if sec.first <= q <= sec.last]
            per_section[sec.name] = score([answers[n] for n in part], [key[n] for n in part],
                                          marks=[marks[n] for n in part])
        per_topic: dict[str, tuple[int, int]] = {}
        for q, state in zip(qs, counts):
            if q in (topics or {}) and state != "dropped":
                right, counted = per_topic.get(topics[q], (0, 0))
                per_topic[topics[q]] = (right + (state == "right"), counted + 1)
        omr = score(answers, key, marks=marks)
        extra = written.get(p.result.roll_number) if written is not None else None
        out.append(Graded(p, key, omr + (extra or 0), counts.count("right"), counts.count("wrong"),
                          counts.count("blank"), to_check=i in flagged, student=students.get(p.result.roll_number),
                          sections=per_section, omr=omr, written=extra, topics=per_topic))
    out.sort(key=lambda g: -g.score)
    for g, (r, pct) in zip(out, rank([g.score for g in out])):
        g.rank, g.percentile = r, pct
    return out


def item_analysis(t: Template, graded: list[Graded]) -> list[list]:
    """One row per question: key, % right, % wrong, % blank, the most chosen wrong answer, and discrimination.

    Discrimination is the share right in the top 27% of scores minus the share right in the bottom 27%:
    near 0 or negative means the question doesn't separate strong from weak students, worth a second look.
    """
    qs = questions(t)
    ordered = sorted(graded, key=lambda g: -g.score)
    group = max(1, round(0.27 * len(ordered)))
    top, bottom = ordered[:group], ordered[-group:]
    rows = []
    for n, q in enumerate(qs):
        counted = [g for g in graded if g.key[n] != DROPPED]
        keys = sorted({g.key[n] for g in graded})
        key = keys[0] if len(keys) == 1 else "by set"
        if not counted:
            rows.append([f"Q{q}", "dropped", None, None, None, "", None])
            continue
        states = Counter(status(g.page.result.answers[q], g.key[n]) for g in counted)
        wrong = Counter(g.page.result.answers[q] for g in counted if status(g.page.result.answers[q], g.key[n]) == "wrong")
        pct = lambda k: round(100 * states[k] / len(counted), 1)
        right_share = lambda gs: sum(status(g.page.result.answers[q], g.key[n]) == "right" for g in gs
                                     if g.key[n] != DROPPED) / max(1, sum(g.key[n] != DROPPED for g in gs))
        common = wrong.most_common(1)
        rows.append([f"Q{q}", key, pct("right"), pct("wrong"), pct("blank"),
                     f"{common[0][0]} ({common[0][1]})" if common else "", round(right_share(top) - right_share(bottom), 2)])
    return rows


def has_written(graded: list[Graded]) -> bool:
    """Whether any student has written marks, so results show the sheet and written parts apart."""
    return any(g.written is not None for g in graded)


ANSWERS = "Answers"  # one entry in the column list, written as a column per question
FIXED = ["Rank", "Roll number", "Name", "Class", "Set", "Score", "Sheet score", "Written", "Right", "Wrong", "Blank",
         "Percentile", "Sheet no.", "File", "Page", ANSWERS, "Review"]


def section_column(name: str) -> str:
    """A section's column name, kept apart from the fixed columns a section could be named after."""
    return f"{name} (section)" if name in FIXED else name


def csv_columns(t: Template, sections: list[str], written: bool) -> list[str]:
    """Every column a results CSV can have for this exam, in the default order."""
    skip = ({"Sheet score", "Written"} if not written else set()) | ({"Sheet no."} if t.serial is None else set())
    cols = [c for c in FIXED if c not in skip]
    at = cols.index("Right")
    return cols[:at] + [section_column(n) for n in sections] + cols[at:]


def write_csv(path: str | Path, t: Template, pages: list[StoredPage], graded: list[Graded],
              students: dict[str, Student], columns: list[str]):
    """The chosen columns in the chosen order: scored pages best first, then pages that couldn't be scored."""
    qs = questions(t)
    reasons: dict[int, list[str]] = {}
    for i, item, why in to_review(pages, students):
        reasons.setdefault(i, []).append(f"{item}: {why}")
    by_page = {g.page.id: g for g in graded}
    index = {p.id: i for i, p in enumerate(pages)}
    order = [g.page for g in graded] + [p for p in pages if p.id not in by_page]
    num = lambda v: "" if v is None else f"{v:g}"
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # the BOM lets Excel open it as UTF-8
        w = csv.writer(f)
        w.writerow([h for c in columns for h in ([f"Q{q}" for q in qs] if c == ANSWERS else [c])])
        for p in order:
            r, g = p.result, by_page.get(p.id)
            student = students.get(r.roll_number) if r else None
            values = {"Rank": g and g.rank, "Roll number": r and r.roll_number, "Name": student and student.name,
                      "Class": student and student.class_name, "Set": r and r.set_code,
                      "Score": num(g and g.score), "Sheet score": num(g and g.omr), "Written": num(g and g.written),
                      "Right": g and g.right, "Wrong": g and g.wrong, "Blank": g and g.blank,
                      "Percentile": num(g and round(g.percentile, 1)), "Sheet no.": r and r.serial,
                      "File": p.file, "Page": p.page, "Review": "; ".join(reasons.get(index[p.id], [])),
                      **{section_column(n): num(v) for n, v in (g.sections.items() if g else ())}}
            row = []
            for c in columns:
                row += [r.answers[q] if r else "" for q in qs] if c == ANSWERS else ["" if values.get(c) is None else values[c]]
            w.writerow(row)


def write_excel(path: str | Path, exam_name: str, t: Template, graded: list[Graded]):
    """Merit list on the first sheet, every answer (green right, red wrong) on the second, item analysis on the third."""
    qs = questions(t)
    wb = Workbook()
    ws = wb.active
    ws.title = "Merit list"
    ws.append([exam_name])
    ws["A1"].font = Font(bold=True, size=14)
    section_names = list(graded[0].sections) if graded else []
    split = ["Sheet", "Written"] if has_written(graded) else []
    numbered = ["Sheet no."] if t.serial else []
    head = ["Rank", "Roll number", "Name", "Class", "Set", "Score", *split, *section_names, "Right", "Wrong", "Blank",
            "Percentile", "To check", "File", "Page", *numbered]
    ws.append(head)
    for g in graded:
        r = g.page.result
        s = g.student
        ws.append([g.rank, r.roll_number, s.name if s else "", s.class_name if s else "", r.set_code, g.score,
                   *([g.omr, g.written] if split else []), *g.sections.values(), g.right, g.wrong, g.blank, round(g.percentile, 1),
                   "yes" if g.to_check else "", Path(g.page.file).name, g.page.page, *(r.serial for _ in numbered)])
    answers = wb.create_sheet("Answers")
    answers.append(["Roll number", "Name", "Set", *(f"Q{q}" for q in qs)])
    fills = {"right": PatternFill("solid", fgColor="C6EFCE"), "wrong": PatternFill("solid", fgColor="FFC7CE")}
    for g in graded:
        r = g.page.result
        answers.append([r.roll_number, g.student.name if g.student else "", r.set_code, *(r.answers[q] for q in qs)])
        for col, (q, k) in enumerate(zip(qs, g.key), start=4):
            fill = fills.get(status(r.answers[q], k))
            if fill:
                answers.cell(answers.max_row, col).fill = fill
    items = wb.create_sheet("Item analysis")
    items.append(["Question", "Key", "% right", "% wrong", "% blank", "Most chosen wrong answer", "Discrimination"])
    for row in item_analysis(t, graded):
        items.append(row)
    sheets = [(ws, 2), (answers, 1), (items, 1)]
    by_set = {g.page.result.set_code: g.key for g in graded}
    changed = [[set_code, f"Q{q}", change] for set_code, key in sorted(by_set.items())
               for q, change in changed_questions(key, qs)]
    if changed:
        ws["C1"] = "The answer key was changed for some questions: see the Key changes sheet."
        notes = wb.create_sheet("Key changes")
        notes.append(["Set", "Question", "Change"])
        for row in changed:
            notes.append(row)
        sheets.append((notes, 1))
    names = list(dict.fromkeys(t for g in graded for t in g.topics))
    if names:
        by_topic = wb.create_sheet("Topics")
        by_topic.append(["Roll number", "Name", *(f"{t} % right" for t in names)])
        for g in graded:
            by_topic.append([g.page.result.roll_number, g.student.name if g.student else "",
                             *(_share(g.topics.get(t)) for t in names)])
        totals = [[sum(g.topics.get(t, (0, 0))[k] for g in graded) for k in (0, 1)] for t in names]
        by_topic.append(["Everyone", "", *(_share(total) for total in totals)])
        by_topic.cell(by_topic.max_row, 1).font = Font(bold=True)
        sheets.append((by_topic, 1))
    for sheet, header_row in sheets:
        for cell in sheet[header_row]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = sheet.cell(header_row + 1, 1)
        sheet.auto_filter.ref = f"A{header_row}:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
        for column in sheet.columns:
            width = max(len(str(c.value or "")) for c in column[header_row - 1:]) + 2
            sheet.column_dimensions[column[0].column_letter].width = min(max(width, 6), 40)
    wb.save(path)


def _share(right_counted: tuple[int, int] | None) -> float | None:
    """Percent right, or None for a topic with no counted question."""
    return round(100 * right_counted[0] / right_counted[1], 1) if right_counted and right_counted[1] else None


def write_scorecards(path: str | Path, exam_name: str, t: Template, graded: list[Graded]):
    """One A4 page per student: score, rank, and every answer next to the key."""
    qs = questions(t)
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setTitle(f"{exam_name} scorecards")
    w, h = A4
    per_column = 50
    marks = {"right": ("right", colors.HexColor("#1f9d55")), "wrong": ("wrong", colors.HexColor("#d64545")),
             "blank": ("blank", colors.grey), "dropped": ("dropped", colors.grey)}
    for g in sorted(graded, key=lambda g: g.page.result.roll_number):
        r = g.page.result
        c.setFont("Helvetica-Bold", 16)
        c.drawString(18 * mm, h - 20 * mm, exam_name)
        c.setFont("Helvetica", 11)
        s = g.student
        who = f"{s.name}    " + (f"Class {s.class_name}    " if s.class_name else "") if s else ""
        c.drawString(18 * mm, h - 28 * mm, f"{who}Roll number {r.roll_number}    Set {r.set_code}")
        c.setFont("Helvetica-Bold", 13)
        c.drawString(18 * mm, h - 38 * mm, f"Score {g.score:g}    Rank {g.rank} of {len(graded)}    "
                                           f"Percentile {g.percentile:.1f}")
        c.setFont("Helvetica", 10)
        parts = (f"Sheet {g.omr:g}    Written {g.written:g}    " if g.written is not None
                 else "Sheet {:g}    Written not given    ".format(g.omr) if has_written(graded) else "")
        c.drawString(18 * mm, h - 45 * mm, f"{parts}Right {g.right}    Wrong {g.wrong}    Blank {g.blank}    "
                     + "    ".join(f"{name} {v:g}" for name, v in g.sections.items()))
        top, row_h = h - 60 * mm, 4.4 * mm
        if g.to_check:
            c.setFillColor(colors.HexColor("#c27c00"))
            c.drawString(18 * mm, h - 51 * mm, "Some marks on this sheet were still being checked.")
            c.setFillColor(colors.black)
        if g.topics:
            text = "By topic:  " + "  ·  ".join(f"{t} {r}/{n} ({_share((r, n)):g}%)" for t, (r, n) in g.topics.items())
            lines = simpleSplit(text, "Helvetica", 9, w - 36 * mm)
            c.setFont("Helvetica", 9)
            for k, line in enumerate(lines):
                c.drawString(18 * mm, h - (57 + 4.2 * k) * mm, line)
            top -= 4.2 * len(lines) * mm
        changes = changed_questions(g.key, qs)
        notes = simpleSplit("Answer key changed: " + "  ·  ".join(f"Q{q} {c}" for q, c in changes), "Helvetica", 8,
                            w - 36 * mm) if changes else []
        bottom = (21 + 3.6 * len(notes)) * mm
        rows = min(per_column, int((top - bottom) // row_h) + 1)  # topics and notes take room from the answers
        col_w = (w - 36 * mm) / max(1, -(-len(qs) // rows))
        c.setFont("Helvetica", 8)
        for n, (q, k) in enumerate(zip(qs, g.key)):
            x = 18 * mm + (n // rows) * col_w
            y = top - (n % rows) * row_h
            sign, colour = marks[status(r.answers[q], k)]
            c.setFillColor(colors.black)
            c.drawString(x, y, f"Q{q}")
            c.drawString(x + 10 * mm, y, r.answers[q] or "-")
            c.setFillColor(colors.grey)
            c.drawString(x + 19 * mm, y, "bonus" if k == BONUS else k if k != DROPPED else "")
            c.setFillColor(colour)
            c.drawString(x + 28 * mm, y, sign)
        c.setFillColor(colors.HexColor("#c27c00"))
        for k, line in enumerate(notes):
            c.drawString(18 * mm, (16.5 + 3.6 * (len(notes) - 1 - k)) * mm, line)
        c.setFillColor(colors.grey)
        c.drawString(18 * mm, 12 * mm, "Columns: question, answer given, correct answer, result.")
        c.setFillColor(colors.black)
        c.showPage()
    c.save()
