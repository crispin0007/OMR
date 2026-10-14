"""Read a student list from Excel: roll number, name and class (US7, FR-017), and batch, email and phone.

The header row can be anywhere near the top; columns are found by their
header ("roll", "name", "class" or "section", "batch" or "group", "email",
"phone" or "mobile"). Every row that can't be used is reported with its row
number instead of being imported; a contact that doesn't look right is left
out and reported, and the student is imported without it.
"""

import re
from dataclasses import dataclass, replace
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class Student:
    roll_number: str
    name: str
    class_name: str = ""
    batch: str = ""
    email: str = ""
    phone: str = ""
    student_id: str = ""  # the directory student this is, on an exam's list; "" when not linked


ID_HEADERS = ("student id", "reg", "admission", "enrol")  # a column of the institute's own student numbers
EMAIL = re.compile(r"[^@\s]+@[^@\s]+\.[^@\s]+")


def check_email(email: str) -> str:
    """The email trimmed, "" for none. Raises ValueError for one that doesn't look like an email."""
    email = email.strip()
    if email and (len(email) > 254 or not EMAIL.fullmatch(email)):
        raise ValueError(f"email {email!r} doesn't look like an email address")
    return email


def check_phone(phone: str) -> str:
    """The phone number with only a leading + and digits, "" for none. Raises ValueError unless it has 7 to 15
    digits (the international limit) and nothing but spaces, dashes, dots and brackets besides."""
    raw = phone.strip()
    if raw.endswith(".0") and raw[:-2].isdigit():
        raw = raw[:-2]  # Excel stored the number as a number
    if not raw:
        return ""
    digits = re.sub(r"[\s\-.()]", "", raw)
    plus = digits.startswith("+")
    if not digits.lstrip("+").isdigit() or digits.count("+") > int(plus) or not 7 <= len(digits) - plus <= 15:
        raise ValueError(f"phone {raw!r} isn't a phone number")
    return digits


def read_students(path: str | Path, digits: int | None) -> tuple[list[Student], list[str]]:
    """Students from the first sheet, and one line per skipped row saying why.

    `digits` is how many digits the sheet's roll number has; Excel drops leading
    zeros from numbers, so numeric roll numbers are padded back to that length.
    None is for sheets identified by barcode: any roll number without spaces.
    Raises ValueError when the file can't be opened or has no roll and name columns.
    """
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises many types for files that aren't Excel workbooks
        raise ValueError(f"can't open {Path(path).name} as an Excel file: {e}") from e
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    for h, header in enumerate(rows[:20]):
        names = [str(c or "").strip().lower() for c in header]
        roll = next((i for i, n in enumerate(names) if ("roll" in n or "symbol" in n)
                     and not any(w in n for w in ID_HEADERS)), None)  # "enrollment" isn't a roll number
        name = next((i for i, n in enumerate(names) if "name" in n), None)
        cls = next((i for i, n in enumerate(names) if "class" in n or "section" in n), None)
        batch = next((i for i, n in enumerate(names) if "batch" in n or "group" in n), None)
        mail = next((i for i, n in enumerate(names) if "mail" in n), None)
        phone = next((i for i, n in enumerate(names) if any(w in n for w in ("phone", "mobile", "contact"))), None)
        sid = next((i for i, n in enumerate(names) if n == "id" or any(w in n for w in ID_HEADERS)), None)
        roll = sid if roll is None else roll  # a directory list may have only the students' IDs
        if roll is not None and name is not None:
            break
    else:
        raise ValueError("no header row with a roll number column and a name column in the first 20 rows")

    students, problems, seen = [], [], {}
    for n, row in enumerate(rows[h + 1:], start=h + 2):
        cell = lambda i: row[i] if i is not None and i < len(row) else None
        raw, who = cell(roll), str(cell(name) or "").strip()
        if raw is None and not who:
            continue  # empty row
        if isinstance(raw, float) and raw.is_integer():
            raw = int(raw)
        number = str(raw).zfill(digits or 0) if isinstance(raw, int) else str(raw or "").strip()
        if not number:
            problems.append(f"row {n}: no roll number")
        elif digits is None and any(c.isspace() for c in number):
            problems.append(f"row {n}: roll number {number!r} has spaces")
        elif digits is not None and (not number.isdigit() or len(number) != digits):
            problems.append(f"row {n}: roll number {number!r} isn't {digits} digits")
        elif not who:
            problems.append(f"row {n}: no name for roll number {number}")
        elif number in seen:
            problems.append(f"row {n}: roll number {number} is also on row {seen[number]}")
        else:
            seen[number] = n
            contact = {}
            for key, i, fix in (("email", mail, check_email), ("phone", phone, check_phone)):
                value = cell(i)
                value = str(int(value)) if isinstance(value, float) and value.is_integer() else str(value or "")
                try:
                    contact[key] = fix(value)
                except ValueError as e:
                    contact[key] = ""
                    problems.append(f"row {n}: {e}; {number} imported without it")
            given = cell(sid)
            given = str(int(given)) if isinstance(given, float) and given.is_integer() else str(given or "").strip()
            if any(c.isspace() for c in given):
                problems.append(f"row {n}: student ID {given!r} has spaces; {number} imported without it")
                given = ""
            students.append(Student(number, who, str(cell(cls) or "").strip(), str(cell(batch) or "").strip(),
                                    **contact, student_id=given))
    return students, problems


def absent(students: dict[str, Student], read_rolls: set[str]) -> list[Student]:
    """Students on the list with no sheet read for their roll number, by roll number."""
    return sorted((s for r, s in students.items() if r not in read_rolls), key=lambda s: s.roll_number)


def same_roll(roll: str) -> str:
    """A roll number as compared between a sheet and the directory: a sheet's grid keeps leading zeros that
    Excel drops, so "00123" and "123" are the same student."""
    return roll.lstrip("0") or "0" if roll.isdigit() else roll


def with_directory(listed: dict[str, Student], directory: list[Student], rolls: set[str]) -> dict[str, Student]:
    """The exam's own students, plus each roll number read that the exam's list doesn't have but the directory
    does, so a sheet shows its student's name without a list made for every exam."""
    known = {same_roll(s.roll_number): s for s in directory}
    out = dict(listed)
    for roll in rolls:
        if roll not in out and (s := known.get(same_roll(roll))):
            out[roll] = s
    return out


def name_key(name: str) -> str:
    """A name as compared between lists: case and spacing don't matter."""
    return " ".join(name.split()).casefold()


def link(students: list[Student], directory: list[Student]) -> tuple[list[Student], list[Student], list[Student]]:
    """An exam's students linked to the directory: (every student with student_id set where it can be, students to
    add to the directory, students left unlinked because several directory students have their name).

    A student with an ID keeps it (a new ID is a new directory student); else exactly one directory student with
    the same name links, the phone number choosing between namesakes; else they are new. New students get "" as
    their ID here, for the store to number."""
    ids = {same_roll(d.roll_number): d.roll_number for d in directory}
    by_name: dict[str, list[Student]] = {}
    for d in directory:
        by_name.setdefault(name_key(d.name), []).append(d)
    out, new, unclear = [], [], []
    for s in students:
        if s.student_id:
            known = ids.get(same_roll(s.student_id))
            s = replace(s, student_id=known or s.student_id)
            if not known:
                ids[same_roll(s.student_id)] = s.student_id
                new.append(s)
        else:
            same = by_name.get(name_key(s.name), [])
            if len(same) > 1 and s.phone:
                same = [d for d in same if d.phone == s.phone] or same
            if len(same) == 1:
                s = replace(s, student_id=same[0].roll_number)
            elif same:
                unclear.append(s)
            else:
                new.append(s)
        out.append(s)
    return out, new, unclear


def for_sheet(students: list[Student], digits: int | None) -> tuple[list[Student], list[str]]:
    """Directory students for an exam whose sheet has a `digits`-digit roll number (None for a barcode): numeric
    roll numbers padded with the zeros Excel drops, and the ones that can't be on the sheet, with why."""
    fit, problems = [], []
    for s in students:
        roll = s.roll_number
        if digits is not None:
            if not roll.isdigit() or len(roll.lstrip("0")) > digits:
                problems.append(f"{roll} {s.name}: not a roll number of {digits} digits")
                continue
            roll = roll.zfill(digits)[-digits:]
        fit.append(replace(s, roll_number=roll, student_id=s.student_id or s.roll_number))
    return fit, problems
