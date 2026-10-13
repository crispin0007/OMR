from pathlib import Path

import pypdfium2
from openpyxl import load_workbook

from omr_core.export import grade, item_analysis, write_excel, write_scorecards
from omr_core.grading import Section
from omr_core.reader import PageResult
from omr_core.store import Store

TEMPLATE = (Path(__file__).parent.parent / "templates" / "test-100.json").read_text(encoding="utf-8")


def exam_with_pages(tmp_path):
    s = Store(tmp_path / "omr.db")
    exam = s.create_exam("Mock 1", TEMPLATE)
    t = s.exam(exam).template
    key = "AB*-" + "A" * 96
    for roll, first, flags in (("111", "ABCD", []), ("222", "BB", [("Q3", "unclear mark")]), ("333", "ABCD", []),
                               ("444", "", [])):
        answers = {q: "A" for q in range(5, 101)}
        answers.update({q + 1: a for q, a in enumerate(first)})
        answers.update({q: answers.get(q, "") for q in range(1, 5)})
        s.add_page(exam, "b", f"{roll}.png", 1, PageResult(roll, "A", answers, flags))
    s.add_page(exam, "b", "bad.png", 1, None, "unreadable: registration mark not found")
    return t, s.pages(exam), key


def test_grade_ranks_ties_bonus_and_dropped(tmp_path):
    t, pages, key = exam_with_pages(tmp_path)
    graded = grade(t, pages, key, 4, -1, 0)
    assert [(g.page.result.roll_number, g.score, g.rank) for g in graded] == [
        ("111", 396, 1), ("333", 396, 1), ("222", 391, 3), ("444", 388, 4)]
    by_roll = {g.page.result.roll_number: g for g in graded}
    # 111: Q1 right, Q2 right, Q3 bonus, Q4 dropped, 96 right -> 99 right
    assert (by_roll["111"].right, by_roll["111"].wrong, by_roll["111"].blank, by_roll["111"].score) == (99, 0, 0, 396)
    # 222: Q1 wrong (B), Q2 right, Q3 bonus even though blank, Q4 dropped
    assert (by_roll["222"].right, by_roll["222"].wrong, by_roll["222"].score) == (98, 1, 391)
    # 444: Q1, Q2 blank, Q3 bonus
    assert (by_roll["444"].blank, by_roll["444"].score) == (2, 388)
    assert by_roll["222"].to_check and not by_roll["111"].to_check
    assert [g.rank for g in graded] == [1, 1, 3, 4] and len(graded) == 4  # the unreadable page is left out
    assert graded[0].percentile == 100


def test_merit_list_and_scorecards(tmp_path):
    t, pages, key = exam_with_pages(tmp_path)
    graded = grade(t, pages, key, 4, -1, 0)
    write_excel(tmp_path / "merit.xlsx", "Mock 1", t, graded)
    wb = load_workbook(tmp_path / "merit.xlsx")
    rows = list(wb["Merit list"].iter_rows(min_row=3, values_only=True))
    assert [(r[0], r[1], r[5]) for r in rows] == [(g.rank, g.page.result.roll_number, g.score) for g in graded]
    answers = wb["Answers"]
    assert answers["A2"].value == graded[0].page.result.roll_number and answers.max_column == 103
    write_scorecards(tmp_path / "cards.pdf", "Mock 1", t, graded)
    pdf = pypdfium2.PdfDocument(tmp_path / "cards.pdf")
    assert len(pdf) == 4
    text = pdf[0].get_textpage().get_text_range()
    assert "Roll number 111" in text and "Rank 1 of 4" in text


def test_sections_and_item_analysis(tmp_path):
    t, pages, key = exam_with_pages(tmp_path)
    sections = [Section("Part 1", 1, 4, 4, -1), Section("Part 2", 5, 100, 1, 0)]
    graded = grade(t, pages, key, 4, -1, 0, sections=sections)
    by_roll = {g.page.result.roll_number: g for g in graded}
    assert by_roll["111"].sections == {"Part 1": 12, "Part 2": 96} and by_roll["111"].score == 108
    assert by_roll["222"].sections == {"Part 1": 7, "Part 2": 96}  # Q1 wrong -1, Q2 right, Q3 bonus

    rows = {r[0]: r for r in item_analysis(t, graded)}
    assert rows["Q1"][1:6] == ["A", 50.0, 25.0, 25.0, "B (1)"]  # 111, 333 right; 222 chose B; 444 blank
    assert rows["Q4"][1] == "dropped"
    assert rows["Q5"][2] == 100.0 and rows["Q5"][6] == 0  # everyone right: separates nobody

    write_excel(tmp_path / "merit.xlsx", "Mock 1", t, graded)
    wb = load_workbook(tmp_path / "merit.xlsx")
    assert [c.value for c in wb["Merit list"][2]][5:8] == ["Score", "Part 1", "Part 2"]
    assert wb["Item analysis"].max_row == 101
    write_scorecards(tmp_path / "cards.pdf", "Mock 1", t, graded)
    assert "Part 1 12" in pypdfium2.PdfDocument(tmp_path / "cards.pdf")[0].get_textpage().get_text_range()


def test_written_marks_add_to_the_score_and_rank(tmp_path):
    t, pages, key = exam_with_pages(tmp_path)
    graded = grade(t, pages, key, 4, -1, 0, written={"444": 20, "222": 4.5})
    assert [(g.page.result.roll_number, g.omr, g.written, g.score, g.rank) for g in graded] == [
        ("444", 388, 20, 408, 1), ("111", 396, None, 396, 2), ("333", 396, None, 396, 2), ("222", 391, 4.5, 395.5, 4)]
    write_excel(tmp_path / "merit.xlsx", "Mock 1", t, graded)
    ws = load_workbook(tmp_path / "merit.xlsx")["Merit list"]
    assert [c.value for c in ws[2]][5:8] == ["Score", "Sheet", "Written"]
    assert [c.value for c in ws[3]][5:8] == [408, 388, 20]
    write_scorecards(tmp_path / "cards.pdf", "Mock 1", t, graded)
    text = pypdfium2.PdfDocument(tmp_path / "cards.pdf")[0].get_textpage().get_text_range()
    assert "Sheet 396" in text and "Written not given" in text  # roll 111 comes first on the cards


def test_written_marks_are_kept_checked_and_backed_up(tmp_path):
    import pytest
    s = Store(tmp_path / "omr.db")
    exam = s.create_exam("Mock", TEMPLATE)
    with pytest.raises(ValueError):
        s.set_written(exam, "111", 5)  # no written part yet
    s.set_written_out_of(exam, 20)
    s.set_written(exam, "111", 12.5)
    with pytest.raises(ValueError):
        s.set_written(exam, "222", 21)
    with pytest.raises(ValueError):
        s.set_written_out_of(exam, 10)  # below a mark already typed
    s.set_written(exam, "222", None)
    assert s.written(exam) == {"111": 12.5} and Store(tmp_path / "omr.db").exam(exam).written_out_of == 20
    s.backup(tmp_path / "backup.db")
    other = Store(tmp_path / "other" / "omr.db")
    other.restore(tmp_path / "backup.db")
    new_id = other.exams()[0][0]
    assert other.written(new_id) == {"111": 12.5} and other.exam(new_id).written_out_of == 20


def test_topics_per_student_on_the_scorecard_and_in_excel(tmp_path):
    t, pages, key = exam_with_pages(tmp_path)
    topics = {1: "Kinematics", 2: "Kinematics", 4: "Kinematics", **{q: "Optics" for q in range(5, 101)}}
    graded = grade(t, pages, key, 4, -1, 0, topics=topics)
    by_roll = {g.page.result.roll_number: g for g in graded}
    # Q4 is dropped, so Kinematics counts Q1 and Q2; Q3 has no topic
    assert by_roll["111"].topics == {"Kinematics": (2, 2), "Optics": (96, 96)}
    assert by_roll["222"].topics == {"Kinematics": (1, 2), "Optics": (96, 96)}
    assert grade(t, pages, key, 4, -1, 0)[0].topics == {}

    write_excel(tmp_path / "merit.xlsx", "Mock 1", t, graded)
    rows = list(load_workbook(tmp_path / "merit.xlsx")["Topics"].iter_rows(values_only=True))
    assert rows[0] == ("Roll number", "Name", "Kinematics % right", "Optics % right")
    assert ("222", None, 50, 100) in rows and rows[-1] == ("Everyone", None, 62.5, 100)  # Excel keeps "" as empty
    write_scorecards(tmp_path / "cards.pdf", "Mock 1", t, graded)
    pdf = pypdfium2.PdfDocument(tmp_path / "cards.pdf")
    assert "Kinematics 1/2 (50%)" in pdf[1].get_textpage().get_text_range()  # cards go by roll number: 222 second
    assert "Q100" in pdf[1].get_textpage().get_text_range()  # every answer still fits on the page


def test_csv_has_the_chosen_columns_in_order(tmp_path):
    import csv

    from omr_core.export import ANSWERS, csv_columns, write_csv
    from omr_core.students import Student

    t, pages, key = exam_with_pages(tmp_path)
    assert csv_columns(t, ["Physics", "Score"], written=False)[6:10] == ["Physics", "Score (section)", "Right", "Wrong"]
    assert "Written" in csv_columns(t, [], written=True) and "Sheet no." not in csv_columns(t, [], written=False)
    graded = grade(t, pages, key, 4, -1, 0, sections=[Section("Score", 1, 4, 1, 0)])
    students = {"111": Student("111", "Asha", "10A")}
    write_csv(tmp_path / "r.csv", t, pages, graded, students,
              ["Name", "Roll number", "Wrong", "Score (section)", ANSWERS, "Review"])
    rows = list(csv.reader(open(tmp_path / "r.csv", encoding="utf-8-sig")))
    assert rows[0][:4] == ["Name", "Roll number", "Wrong", "Score (section)"] and rows[0][4:6] == ["Q1", "Q2"]
    assert len(rows[0]) == 4 + 100 + 1
    assert rows[1][:4] == ["Asha", "111", "0", "3"]  # best first (Q3 is a bonus); 0 wrong is written, not left empty
    assert [r[1] for r in rows[1:]] == ["111", "333", "222", "444", ""]  # the unreadable page comes last
    assert rows[-1][0] == "" and rows[-1][-1].startswith("page: unreadable") and rows[-1][4] == ""
    assert "Q3: unclear mark" in rows[3][-1] and "not on the student list" in rows[3][-1]
