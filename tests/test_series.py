from pathlib import Path

import pytest
from openpyxl import load_workbook

from omr_core.reader import PageResult
from omr_core.series import series_results, write_series_excel
from omr_core.store import Store
from omr_core.students import Student

TEMPLATE = (Path(__file__).parent.parent / "templates" / "test-100.json").read_text(encoding="utf-8")


def sheet(roll: str, right: int) -> PageResult:
    """A sheet with the first `right` questions right (A) and the rest wrong (B)."""
    return PageResult(roll, "A", {q: "A" if q <= right else "B" for q in range(1, 101)}, [])


def test_series_ranks_students_on_their_combined_score(tmp_path):
    s = Store(tmp_path / "omr.db")
    one = s.create_exam("Mock 1", TEMPLATE, {"series": "JEE"})
    s.set_students(one, [Student("111", "Asha", "12"), Student("222", "Bina", "12"), Student("333", "Chet", "11")])
    two = s.create_exam("Mock 2", TEMPLATE, {"series": "JEE"})
    s.create_exam("Elsewhere", TEMPLATE, {"series": "NEET"})
    for exam, results in ((one, [("111", 60), ("222", 80), ("333", 50)]), (two, [("111", 90), ("222", 40),
                                                                                 ("111", 10), ("9?9", 100)])):
        s.save_scheme(exam, "A" * 100, 1, 0, 0)
        for roll, right in results:
            s.add_page(exam, "b", f"{roll}.png", 1, sheet(roll, right))
    exams, rows, problems = series_results(s, "JEE")
    assert exams == ["Mock 1", "Mock 2"] and problems == []
    # 111 read twice in Mock 2 keeps the better sheet; 9?9 belongs to nobody until Review fixes it;
    # 333 missed Mock 2 and scores nothing for it
    assert [(r.rank, r.roll_number, r.name, r.scores, r.total, r.taken, r.average) for r in rows] == [
        (1, "111", "Asha", [60, 90], 150, 2, 75.0), (2, "222", "Bina", [80, 40], 120, 2, 60.0),
        (3, "333", "Chet", [50, None], 50, 1, 50.0)]

    write_series_excel(tmp_path / "series.xlsx", "JEE", exams, rows)
    ws = load_workbook(tmp_path / "series.xlsx")["Series merit list"]
    assert [c.value for c in ws[2]] == ["Rank", "Roll number", "Name", "Class", "Mock 1", "Mock 2", "Total",
                                        "Tests taken", "Average %", "Movement", "Improvement %"]
    assert [c.value for c in ws[5]] == [3, "333", "Chet", "11", 50, "absent", 50, 1, 50, 0, None]


def test_an_exam_without_a_key_is_named_not_counted(tmp_path):
    s = Store(tmp_path / "omr.db")
    exam = s.create_exam("Mock 1", TEMPLATE, {"series": "JEE"})
    s.add_page(exam, "b", "1.png", 1, sheet("111", 10))
    assert series_results(s, "JEE")[2] == ["Mock 1: no results yet (no answer key, or no sheets read)"]
    s.save_scheme(exam, "AE", 1, 0, 0)
    assert "Mock 1: line 1: E is not an option" in series_results(s, "JEE")[2][0]


def test_weights_best_of_movement_and_strengths(tmp_path):
    from omr_core.grading import Section
    from omr_core.series import exam_summary
    s = Store(tmp_path / "omr.db")
    exams = [s.create_exam(f"Mock {n}", TEMPLATE, {"series": "CEE"}) for n in (1, 2, 3)]
    s.set_students(exams[0], [Student("111", "Asha", "12", "Morning"), Student("222", "Bina", "12", "Day")])
    results = [[("111", 50), ("222", 70)], [("111", 60), ("222", 40)], [("111", 90)]]
    for exam, sheets in zip(exams, results):
        s.save_scheme(exam, "A" * 100, 1, 0, 0)
        s.save_sections(exam, [Section("Physics", 1, 50, 1, 0, 0), Section("Chemistry", 51, 100, 1, 0, 0)])
        for roll, right in sheets:
            s.add_page(exam, "b", f"{roll}.png", 1, sheet(roll, right))
    _, rows, _ = series_results(s, "CEE")
    asha, bina = rows
    assert (asha.total, asha.rank, bina.total) == (200, 1, 110)
    assert asha.exam_ranks == [2, 1, 1] and bina.exam_ranks == [1, 2, None] and asha.batch == "Morning"
    # after Mock 2 Asha led 110 to 110, tied first; Mock 3 kept her first and dropped Bina to second
    assert (asha.movement, bina.movement) == (0, -1)
    assert asha.improvement == 40.0 and bina.improvement == -30.0
    assert asha.sections == {"Physics": 100.0, "Chemistry": 33.3}  # 150 of 150 and 50 of 150
    assert exam_summary(["Mock 1", "Mock 2", "Mock 3"], rows)[2] == ("Mock 3", 1, 90.0, 90.0)

    with pytest.raises(ValueError):
        s.save_series_settings("CEE", 1, {exams[2]: 11})
    s.save_series_settings("CEE", 2, {exams[2]: 2})  # the final mock counts double; best 2 of 3
    _, rows, _ = series_results(s, "CEE")
    asha = next(r for r in rows if r.roll_number == "111")
    assert asha.counted == [False, True, True] and asha.total == 60 + 2 * 90
    assert asha.average == round((60 + 2 * 90) / 3, 1)
    assert s.series_settings("CEE")["best_of"] == 2


def test_a_student_under_a_new_roll_number_in_each_test_is_one_row(tmp_path):
    from omr_core.progress import attendance, history

    s = Store(tmp_path / "omr.db")
    s.rolls_change = True
    s.save_student(Student("S0001", "Asha Rai"))
    one = s.create_exam("Mock 1", TEMPLATE, {"series": "JEE"})
    two = s.create_exam("Mock 2", TEMPLATE, {"series": "JEE"})
    s.set_students(one, [Student("105", "Asha Rai", student_id="S0001"), Student("106", "Bina")])
    s.set_students(two, [Student("212", "Asha Rai", student_id="S0001"), Student("105", "Chet")])
    for exam, results in ((one, [("105", 60), ("106", 70)]), (two, [("212", 90), ("105", 40)])):
        s.save_scheme(exam, "A" * 100, 1, 0, 0)
        for roll, right in results:
            s.add_page(exam, "b", f"{roll}.png", 1, sheet(roll, right))
    _, rows, _ = series_results(s, "JEE")
    # Asha is one row under her latest roll number; Chet sat as 105 after Asha did, and is someone else
    assert sorted((r.roll_number, r.name, r.scores) for r in rows) == [
        ("105", "Chet", [None, 40]), ("106", "Bina", [70, None]), ("212", "Asha Rai", [60, 90])]
    assert [e.score for e in history(s, "S0001")] == [60, 90]
    assert attendance(s, ["S0001"], [one, two]) == {"S0001": [True, True]}
    assert s.exams_taken()["S0001"] == 2

    s.rolls_change = False  # roll numbers are Student IDs: an unlinked 105 is the same student in both
    _, rows, _ = series_results(s, "JEE")
    assert len(rows) == 3


def test_a_students_weak_topics_add_up_over_their_exams(tmp_path):
    from omr_core.progress import history, topic_averages

    s = Store(tmp_path / "omr.db")
    s.save_student(Student("111", "Asha"))
    for name, right in (("Mock 1", 10), ("Mock 2", 30)):
        exam = s.create_exam(name, TEMPLATE)
        s.save_scheme(exam, "A" * 100, 1, 0, 0)
        s.set_topics(exam, {**{q: "Kinematics" for q in range(1, 41)}, **{q: "Optics" for q in range(41, 101)}})
        s.add_page(exam, "b", "a.png", 1, sheet("111", right))
    # Kinematics: 10 of 40, then 30 of 40; Optics: none right in either
    assert topic_averages(history(s, "111")) == [("Optics", 0.0, 120, 2), ("Kinematics", 50.0, 80, 2)]
