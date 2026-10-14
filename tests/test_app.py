import os

import cv2
import numpy as np
import pytest

pytest.importorskip("PySide6")
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QCoreApplication, QEvent, QRectF, QSettings, Qt  # noqa: E402
from PySide6.QtWidgets import QApplication, QLabel  # noqa: E402

from omr_app import designer, scanner, theme  # noqa: E402
from omr_app.columns import ColumnsDialog  # noqa: E402
from omr_app.scan import FOLDER  # noqa: E402
from omr_app.window import TEMPLATES, MainWindow  # noqa: E402
from omr_core.reader import SCALE, read_page  # noqa: E402
from omr_core.template import load, parse  # noqa: E402
from sheets import draw, random_marks, scan  # noqa: E402


@pytest.fixture
def app(tmp_path):
    QSettings.setPath(QSettings.IniFormat, QSettings.UserScope, str(tmp_path))  # keep the user's settings out of it
    return QApplication.instance() or QApplication([])


@pytest.fixture(autouse=True)
def close_windows():
    """Close what a test opened while the app still exists: a window left open until Python exits is torn
    down after Qt, its scanner thread or database half gone, and the test run crashed on closing."""
    yield
    app = QApplication.instance()
    if app is None:
        return
    for w in app.topLevelWidgets():
        w.close()  # MainWindow stops its scanner thread and closes the database
        w.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
    app.processEvents()


def window(tmp_path) -> MainWindow:
    w = MainWindow(tmp_path / "omr.db")
    w.open_exam(w.store.create_exam("Mock", (TEMPLATES / "test-100.json").read_text(encoding="utf-8")))
    return w


def sheet_files(tmp_path, seeds):
    t = load(TEMPLATES / "test-100.json")
    out = []
    for seed in seeds:
        marks, answers, roll, set_code = random_marks(t, seed)
        path = tmp_path / f"sheet{seed}.png"
        cv2.imwrite(str(path), scan(draw(t, marks), seed=seed))
        out.append((str(path), answers, roll, set_code))
    return out


def test_reads_sheets_and_scores_them(app, tmp_path):
    from PySide6.QtTest import QTest
    (path, answers, roll, set_code), = sheet_files(tmp_path, [3])
    w = window(tmp_path)
    c = w.check
    c.read_files([path])
    cell = lambda col: c.table.item(0, col).text()
    assert (cell(2), cell(4), cell(6)) == (roll, set_code, "")  # no key yet, no score
    key = [answers[q] or "A" for q in c.questions]
    c.key_edit.setPlainText(f"{set_code}: " + "".join(key))
    right = sum(1 for q in c.questions if answers[q])
    assert cell(6) == str(right)
    c.wrong.setValue(-1)
    assert cell(6) == str(right)  # blanks aren't wrong
    c.wrong.setValue(0.25)
    assert c.wrong.value() == 0  # a stored bonus for wrong answers, never meant, loads as no negative marking
    c.wrong.selectAll()
    QTest.keyClicks(c.wrong, "0.25")  # people type the marks taken off; the box shows the minus
    c.wrong.interpretText()
    assert c.wrong.value() == -0.25 and c.wrong.text() == "−0.25"
    c.table.selectRow(0)
    assert c.review_list.item(0).text() == "Nothing to check"
    assert w.dashboard.read.value.text() == "1" and w.dashboard.average.value.text() == f"{right:.1f}"
    assert w.results.table.item(0, 1).text() == roll
    assert w.results.table.item(0, 0).text() == "1"
    w.show()
    for page in ["Dashboard", "Exams", "Templates", "Map printed sheet", "Settings", *w.exam_page.names]:  # all paint
        w.go(page)
        assert not w.grab().isNull()


def test_dashboard_opens_the_sheets_behind_it(app, tmp_path):
    files = sheet_files(tmp_path, [1, 2])
    (tmp_path / "blank.png").write_bytes(cv2.imencode(".png", 255 + 0 * cv2.imread(files[0][0]))[1].tobytes())
    w = window(tmp_path)
    w.check.read_files([f[0] for f in files] + [str(tmp_path / "blank.png")])  # a page with no marks can't be read
    assert w.check.to_check() == [2]
    assert w.dashboard.to_check.value.text() == "1" and w.dashboard.batches.value.text() == "1"
    w.dashboard.to_check.mousePressEvent(None)
    assert w.showing() is w.review
    w.dashboard.open_clean()
    assert w.showing() is w.check
    assert [w.check.table.isRowHidden(r) for r in range(3)] == [False, False, True]
    w.check.show_sheets(None)
    assert not any(w.check.table.isRowHidden(r) for r in range(3))


def test_exam_is_there_after_reopening(app, tmp_path):
    (path, answers, roll, set_code), = sheet_files(tmp_path, [4])
    w = window(tmp_path)
    w.check.read_files([path])
    w.check.key_edit.setPlainText("".join(answers[q] or "A" for q in w.check.questions))
    w.check.wrong.setValue(-0.25)
    w.close()

    w = MainWindow(tmp_path / "omr.db")  # opens the last exam
    c = w.check
    assert c.exam.name == "Mock" and w.showing() is w.dashboard
    assert c.table.item(0, 2).text() == roll and c.wrong.value() == -0.25
    assert c.table.item(0, 6).text() == str(sum(1 for q in c.questions if answers[q]))
    c.table.selectRow(0)
    assert c.page_view.pixmap() is not None and not c.page_view.pixmap().isNull()  # the page image was saved


def test_first_run_starts_on_exams(app, tmp_path):
    w = MainWindow(tmp_path / "omr.db")
    assert w.check.exam is None and w.showing() is w.exams
    assert not w.check.open_button.isEnabled()


def test_settings_apply_and_are_kept(app, tmp_path):
    w = MainWindow(tmp_path / "omr.db")
    w.settings.theme.setCurrentText("Dark")
    w.settings._set(accent="#123456")
    w.settings.font.setCurrentText("Large")
    assert theme.tokens()["bg"] == theme.DARK["bg"] and theme.tokens()["accent"] == "#123456"
    assert theme.load_prefs() == theme.Prefs("Dark", "#123456", "Large", "Comfortable")
    w.toggle_theme()
    assert theme.load_prefs().theme == "Light" and theme.tokens()["bg"] == theme.LIGHT["bg"]


def test_bad_saved_settings_fall_back_to_defaults(app):
    s = QSettings(QSettings.IniFormat, QSettings.UserScope, "OMR", "OMR")
    for k, v in {"theme": "Neon", "accent": "not a colour", "font": "Huge", "density": 3}.items():
        s.setValue(k, v)
    s.sync()
    assert theme.load_prefs() == theme.Prefs()


def test_review_fixes_a_flagged_answer_and_logs_it(app, tmp_path):
    t = load(TEMPLATES / "test-100.json")
    marks, answers, roll, set_code = random_marks(t, 5)
    q_grid = next(i for i, g in enumerate(t.grids) if g.kind == "questions")
    for c in range(4):
        marks.pop((q_grid, 0, c), None)
    marks[q_grid, 0, 2] = "half"  # question 1 half filled: the reader flags it, never guesses
    path = tmp_path / "half.png"
    cv2.imwrite(str(path), scan(draw(t, marks)))
    w = window(tmp_path)
    c = w.check
    c.read_files([str(path)])
    c.key_edit.setPlainText("".join(["C"] + [answers[q] or "A" for q in c.questions[1:]]))
    before = float(c.table.item(0, 6).text())
    w.go("Review")
    assert [item for _, item, _ in w.review.items] == ["Q1"]
    assert w.review.picture.pixmap() is not None and not w.review.picture.pixmap().isNull()
    w.review.fix("Q1", "C")
    assert w.review.items == [] and float(c.table.item(0, 6).text()) == before + 1
    assert w.store.audit(c.exam.id)[0][4:] == ("Q1", "", "C")
    assert w.review.log.count() == 1


def test_results_export_asks_while_marks_wait_for_review(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    files = sheet_files(tmp_path, [1, 2])
    (tmp_path / "blank.png").write_bytes(cv2.imencode(".png", 255 + 0 * cv2.imread(files[0][0]))[1].tobytes())
    w = window(tmp_path)
    w.check.read_files([f[0] for f in files] + [str(tmp_path / "blank.png")])  # the blank page waits for review
    w.check.key_edit.setPlainText("A" * 100)
    assert w.results.table.rowCount() == 2
    out = tmp_path / "merit.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), ""))
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.No)
    w.results.excel.click()
    assert not out.exists()  # said no
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    w.results.excel.click()
    assert out.exists()


def test_imported_students_show_by_name(app, tmp_path, monkeypatch):
    from openpyxl import Workbook
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    (path, answers, roll, set_code), = sheet_files(tmp_path, [6])
    wb = Workbook()
    wb.active.append(["Roll number", "Name", "Class"])
    wb.active.append([roll, "Asha Rai", "10 A"])
    wb.save(tmp_path / "students.xlsx")
    w = window(tmp_path)
    w.check.read_files([path])
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(tmp_path / "students.xlsx"), ""))
    shown = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: shown.append(a[2]))
    w.students.import_button.click()
    assert shown and shown[0].startswith("Imported 1 student into")
    assert w.check.table.item(0, 3).text() == "Asha Rai" and w.students.table.item(0, 1).text() == "Asha Rai"
    w.check.key_edit.setPlainText("A" * 100)
    assert w.results.table.item(0, 2).text() == "Asha Rai"


def test_designer_maps_a_sheet_and_saves_a_template(app, tmp_path, monkeypatch):
    t = load(TEMPLATES / "test-100.json")
    blank = tmp_path / "blank.png"
    cv2.imwrite(str(blank), scan(draw(t, {}), angle=0.5))
    w = window(tmp_path)
    w.show()
    w.go("Map printed sheet")
    d = w.designer
    assert d.intro.isVisible() and not d.split.isVisible()  # says what the page is for, and points to New exam
    from omr_app import exams
    monkeypatch.setattr(exams.NewExamWizard, "exec", lambda self: False)
    d.new_exam_wanted.emit()
    assert w.showing() is w.exams
    w.go("Map printed sheet")
    d.load_scan(str(blank))
    assert not d.intro.isVisible() and d.split.isVisible()
    s = d.gray.shape[1] / t.page[0]
    d.mark_clicked(170 * s, 80 * s)  # blank paper, nowhere near a mark
    assert d.error.text() and not d.marks
    for x, y in t.markers:
        d.mark_clicked(x * s + 5, y * s - 4)
    assert len(d.marks) == 4 and d.canvas.drag_boxes
    for g in t.grids:
        (x0, y0), (x1, y1) = g.centre(0, 0), g.centre(g.rows - 1, g.cols - 1)
        d.canvas.boxed.emit(QRectF((x0 - 4) * SCALE, (y0 - 4) * SCALE, (x1 - x0 + 8) * SCALE, (y1 - y0 + 8) * SCALE))
    assert [b["kind"] for b in d.blocks] == [g.kind for g in t.grids]
    assert [b.get("first") for b in d.blocks] == [None, None, 1, 26, 51, 76]
    assert d.blocks[1]["options"] == "ABCD"

    marks, answers, roll, set_code = random_marks(t, 5)
    filled = tmp_path / "filled.png"
    cv2.imwrite(str(filled), scan(draw(t, marks), seed=5))
    d.load_scan(str(filled))  # a new scan starts the corners again
    assert not d.blocks
    for x, y in t.markers:
        d.mark_clicked(x * s, y * s)
    for g in t.grids:
        (x0, y0), (x1, y1) = g.centre(0, 0), g.centre(g.rows - 1, g.cols - 1)
        d.canvas.boxed.emit(QRectF((x0 - 4) * SCALE, (y0 - 4) * SCALE, (x1 - x0 + 8) * SCALE, (y1 - y0 + 8) * SCALE))
    d.test_read()
    assert f"roll number {roll}, set {set_code}" in d.test_result.text()

    monkeypatch.setattr(designer, "USER_TEMPLATES", tmp_path / "templates")
    monkeypatch.setattr(designer.QFileDialog, "getSaveFileName", lambda _p, _t, path, _f: (path, ""))
    monkeypatch.setattr(designer.QMessageBox, "information", lambda *a: None)
    d.name.setText("My sheet")
    d.save()
    saved = load(tmp_path / "templates" / "My sheet.json")
    result = read_page(scan(draw(t, marks), seed=6), saved)
    assert (result.roll_number, result.set_code, result.answers) == (roll, set_code, answers)


class FakeScanner:
    name = "Test scanner"

    def __init__(self, pages, problem=""):
        self.pages_, self.problem = pages, problem

    def pages(self, dpi, duplex):
        assert (dpi, duplex) == (200, True)
        yield from self.pages_
        if self.problem:
            raise scanner.ScanStopped(self.problem)


def drain(page):
    """Read every file waiting on the Scan page, as its timer would."""
    while page.reading_files:
        page.step()


def run_scan(app, page):
    page.start()
    page.thread.wait()
    for _ in range(5):
        app.processEvents()


def test_scanning_reads_each_page_and_stops_on_a_jam(app, tmp_path, monkeypatch):
    t = load(TEMPLATES / "test-100.json")
    fronts = [random_marks(t, seed) for seed in (1, 2, 3)]
    back = np.full((2339, 1654, 3), 240, np.uint8)
    pages = [p for m in fronts[:2] for p in (scan(draw(t, m[0]), seed=1), back)]
    monkeypatch.setattr(scanner, "scanners", lambda: [FakeScanner(pages, "Paper jam")])
    w = window(tmp_path)
    s = w.scan
    s.find_scanners()
    assert s.source.currentText() == "Test scanner"
    s.duplex.setChecked(True)
    run_scan(app, s)
    assert s.thread is None and "Paper jam" in s.status.text() and fronts[1][2] in s.status.text()
    assert (s.scanned.value.text(), s.read.value.text(), s.skipped.value.text()) == ("4", "2", "2")
    assert [p.result.roll_number for p in w.store.pages(w.check.exam.id)] == [fronts[0][2], fronts[1][2]]
    assert fronts[1][2] in s.latest_text.text()
    assert w.check.table.rowCount() == 2 and w.check.table.item(1, 2).text() == fronts[1][2]  # shown on stopping

    # The operator clears the jam and feeds the rest.
    monkeypatch.setattr(scanner, "scanners", lambda: [FakeScanner([scan(draw(t, fronts[2][0]))])])
    s.find_scanners()
    run_scan(app, s)
    assert s.status.text().startswith("Feeder empty") and len(w.store.pages(w.check.exam.id)) == 3


def test_scanning_from_a_folder(app, tmp_path, monkeypatch):
    monkeypatch.setattr(scanner, "scanners", lambda: [])
    folder = tmp_path / "scans"
    folder.mkdir()
    files = sheet_files(folder, [5, 6])
    (folder / "broken.png").write_bytes(b"not an image")
    w = window(tmp_path)
    s = w.scan
    s.find_scanners()
    assert s.source.currentText() == FOLDER and "No scanner found" in s.status.text()
    s.folder.setText(str(folder))
    s.start()
    s.poll_folder()  # sizes noted, nothing read yet
    assert s.read.value.text() == "0"
    s.poll_folder()
    drain(s)
    assert s.read.value.text() == "2" and "broken.png" in s.status.text() and s.live.rowCount() == 2
    s.poll_folder()  # nothing read twice, the broken file isn't retried
    drain(s)
    assert len(w.store.pages(w.check.exam.id)) == 2
    assert sorted(p.result.roll_number for p in w.store.pages(w.check.exam.id)) == sorted(f[2] for f in files)
    s.stop()
    assert not s.running


def test_new_exam_wizard_makes_the_exam_and_its_sheet(app, tmp_path, monkeypatch):
    import json
    from omr_app import exams, new_exam
    w = MainWindow(tmp_path / "omr.db")
    logo = tmp_path / "logo.png"
    cv2.imwrite(str(logo), np.full((30, 90, 3), (0, 120, 0), np.uint8))
    monkeypatch.setattr(new_exam.QFileDialog, "getOpenFileName", lambda *a: (str(logo), ""))

    def fill(z):
        assert z.next.text() == "Next" and not z.back.isEnabled()
        z.forward()  # Start: make the sheet for me
        z.forward()
        assert z.error.text() == "Give the exam a name"
        z.name.setText("Weekly test")
        z.class_name.setText("Class 10")
        z.minutes.setValue(60)
        z.forward()
        z.add_section()
        z.add_section()
        z.sections.item(0, 0).setText("Physics")
        z.sections.item(0, 1).setText("20")
        z.sections.item(1, 0).setText("Physics")
        assert z.sections.item(1, 2).text() == "21 to 45" and z.total.value() == 45
        z.forward()
        assert "own name" in z.error.text()
        z.sections.item(1, 0).setText("Chemistry")
        z.forward()
        z.options.setValue(5)
        z.roll.setValue(4)
        z.sets.setCurrentIndex(0)  # no sets
        z.forward()
        assert z.section_marks.rowCount() == 2
        z.section_marks.cellWidget(0, 1).setValue(4)
        z.section_marks.cellWidget(0, 2).setValue(-1)
        z.section_marks.item(1, 3).setCheckState(Qt.Checked)
        z.forward()
        z.colour.setCurrentIndex(z.colour.findText("Blue"))
        z.choose_logo()
        z.forward()
        assert z.next.text() == "Create exam" and "45 questions, 5 options each" in z.summary.text()
        assert "Chemistry: 25 questions, 1 right, 0 wrong, partial marking" in z.summary.text()
        z.forward()
        return z.result()

    monkeypatch.setattr(exams.NewExamWizard, "exec", fill)
    w.exams.new_exam()
    exam = w.check.exam
    assert exam.name == "Weekly test" and w.showing() is w.print_sheet and w.print_sheet.save.isEnabled()
    assert exam.details["class"] == "Class 10" and exam.details["minutes"] == 60
    assert "Class 10" in w.exam_page.about.text() and "60 minutes" in w.exam_page.about.text()
    assert [(s.name, s.first, s.last, s.correct, s.wrong, s.partial) for s in exam.sections] == [
        ("Physics", 1, 20, 4, -1, False), ("Chemistry", 21, 45, 1, 0, True)]
    t = json.loads(exam.template_json)
    assert t["print"]["colour"] == "#1565c0" and t["print"]["logo_jpeg"]
    assert [g.get("multi", False) for g in t["grids"] if g["kind"] == "questions"][-1]

    def own(z):
        z.own_sheet.setChecked(True)
        z.template_path = ""
        z.forward()
        assert "template" in z.error.text()
        z.template_path = str(TEMPLATES / "test-100.json")
        z.forward()
        z.name.setText("Printed one")
        z.forward()
        assert STEP(z) == "Marking" and z.section_marks.isHidden()
        z.correct.setValue(2)
        z.wrong.setValue(-0.5)
        z.forward()
        assert STEP(z) == "Review and create" and "Your printed sheet" in z.summary.text()
        z.forward()
        return z.result()

    STEP = lambda z: new_exam.STEPS[z.steps.currentIndex()]  # noqa: E731
    monkeypatch.setattr(exams.NewExamWizard, "exec", own)
    w.exams.new_exam()
    assert w.check.exam.name == "Printed one" and w.showing() is w.scan
    e = w.check.exam
    assert (e.correct, e.wrong, e.blank, e.details["minutes"]) == (2, -0.5, 0, 90)


def test_new_exam_wizard_true_false_and_layout_problems(app, tmp_path):
    import json
    from omr_app.new_exam import NewExamWizard
    w = MainWindow(tmp_path / "omr.db")
    z = NewExamWizard(w, w.store)
    z.name.setText("Quick quiz")
    z.true_false.setChecked(True)
    z.total.setValue(20)
    z.roll.setValue(0)
    z.update_preview()
    assert z.options.value() == 2 and not z.options.isEnabled() and z.made
    assert {g["options"] for g in z.made["grids"] if g["kind"] == "questions"} == {"TF"}
    z.total.setValue(300)
    z.options.setEnabled(True)
    z.true_false.setChecked(False)
    z.options.setValue(8)
    z.update_preview()
    assert z.made is None and z.layout_problem
    assert z.check_step("Review and create") == z.layout_problem
    z.total.setValue(40)
    z.update_preview()
    z.create()
    assert json.loads(w.store.exam(z.exam_id).template_json)["print"]


def test_answer_key_grid_keyboard_and_csv(app, tmp_path, monkeypatch):
    from PySide6.QtTest import QTest
    from omr_app import answer_key
    w = window(tmp_path)
    k = w.check.key_tab
    g = k.grid
    assert g.rowCount() == 100 and g.horizontalHeaderItem(2).text() == "A" and g.currentRow() == 0
    QTest.keyClick(g, Qt.Key_B)
    QTest.keyClick(g, Qt.Key_D)
    assert g.currentRow() == 2 and w.check.key_edit.toPlainText() == "BD"
    QTest.keyClick(g, Qt.Key_Up)
    QTest.keyClick(g, Qt.Key_A, Qt.ShiftModifier)  # a second correct option for question 2
    assert w.check.key_edit.toPlainText() == "B, AD" and g.currentRow() == 1
    QTest.keyClick(g, Qt.Key_Down)
    QTest.keyClick(g, Qt.Key_Asterisk)
    QTest.keyClick(g, Qt.Key_Return)
    QTest.keyClick(g, Qt.Key_C)
    assert w.check.key_edit.toPlainText() == "B, AD, *, -, C" and "4 of 100" in k.count.text()
    g.setCurrentCell(4, 2)
    g.cellClicked.emit(4, 2)  # click A for question 5
    assert w.check.key_edit.toPlainText() == "B, AD, *, -, A"
    QTest.keyClick(g, Qt.Key_Space)  # the current row (question 5) is dropped again
    g.setCurrentCell(0, 2)
    QTest.keyClick(g, Qt.Key_Delete)
    assert w.check.key_edit.toPlainText() == "-, AD, *"
    assert w.store.exam(w.check.exam.id).answer_key == "-, AD, *"  # saved as it is set

    k.sets.setCurrentText("B")
    assert "uses the key for all sets" in k.count.text()
    QTest.keyClick(g, Qt.Key_C)
    assert w.check.key_edit.toPlainText() == "-, AD, *\nB: -C"

    path = tmp_path / "key.csv"
    monkeypatch.setattr(answer_key.QFileDialog, "getSaveFileName", lambda *a: (str(path), ""))
    k.export_csv()
    assert path.read_text().splitlines()[:3] == ["set,question,answer", ",1,", ",2,AD"]
    w.check.key_edit.clear()
    monkeypatch.setattr(answer_key.QFileDialog, "getOpenFileName", lambda *a: (str(path), ""))
    k.import_csv()
    assert w.check.key_edit.toPlainText() == "-, AD, *\nB: -C"
    path.write_text("question,answer\n1,Z\n")
    warned = []
    monkeypatch.setattr(answer_key.QMessageBox, "warning", lambda *a: warned.append(a[2]))
    k.import_csv()
    assert "Z is not an option" in warned[0] and w.check.key_edit.toPlainText() == "-, AD, *\nB: -C"

    k.paste_button.click()
    assert g.isHidden() and not w.check.key_edit.isHidden()



def test_a_question_with_a_mistake_is_made_bonus_dropped_or_either_and_can_be_undone(app, tmp_path):
    from PySide6.QtTest import QTest
    from omr_app import answer_key
    w = window(tmp_path)
    c, k, g = w.check, w.check.key_tab, w.check.key_tab.grid
    for key in (Qt.Key_B, Qt.Key_D, Qt.Key_A):
        QTest.keyClick(g, key)
    assert c.full_marks() == 3 and w.store.key_log(c.exam.id) == []  # typing the key in isn't a change
    g.setCurrentCell(0, 2)
    QTest.keyClick(g, Qt.Key_C, Qt.ShiftModifier)
    QTest.keyClick(g, Qt.Key_Slash)  # B or C

    def choose(row, text):
        next(a for a in k.mistake_menu(row).actions() if a.text().startswith(text)).trigger()
    choose(1, "Bonus")
    choose(2, "Drop")
    assert c.key_edit.toPlainText() == "B/C, *"  # Q3 dropped at the end of the key: the paper is now 2 questions
    choose(2, "Undo the last change (back to A)")
    assert c.key_edit.toPlainText() == "B/C, *, A" and c.full_marks() == 3
    assert k.changes.text() == ("Changed (1 bonus, 1 with either answer): Q1 B or C accepted  ·  "
                                "Q2 bonus, full marks to everyone")
    assert g.item(1, k.bonus_col).text() == "Bonus" and g.item(0, k.bonus_col).text() == "B or C"

    log = w.store.key_log(c.exam.id)
    assert [r[3:] for r in log] == [(3, "-", "A"), (3, "A", "-"), (2, "D", "*"), (1, "BC", "B/C"), (1, "B", "BC")]
    dialog = answer_key.LogDialog(k)
    dialog.table.selectRow(2)  # Q2 made a bonus
    dialog.undo.click()
    assert c.key_edit.toPlainText() == "B/C, D, A"
    assert w.store.key_log(c.exam.id)[0][3:] == (2, "*", "D")


def test_a_page_that_cant_be_read_says_why_and_offers_mapping(app, tmp_path):
    w = window(tmp_path)
    blank = tmp_path / "blank.png"
    cv2.imwrite(str(blank), np.full((2339, 1654, 3), 255, np.uint8))
    w.check.read_files([str(blank)])
    c = w.check
    assert c.table.item(0, 2).text().startswith("Can't read this page: ") and c.table.columnSpan(0, 2) == 5
    assert "1 of 1 pages couldn't be read" in c.mismatch.text() and not c.map_button.isHidden()
    c.map_button.click()
    assert w.showing() is w.designer


def test_pages_can_be_deleted_or_moved_to_another_exam(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QInputDialog, QMessageBox
    from PySide6.QtCore import QItemSelectionModel
    files = sheet_files(tmp_path, [1, 2, 3])
    w = window(tmp_path)
    other = w.store.create_exam("Retest", (TEMPLATES / "test-100.json").read_text(encoding="utf-8"))
    c = w.check
    c.read_files([f[0] for f in files])
    assert not c.delete_button.isEnabled()
    c.table.selectRow(0)
    c.table.selectionModel().select(c.table.model().index(2, 0),
                                    QItemSelectionModel.Select | QItemSelectionModel.Rows)
    assert c.selected() == [0, 2] and c.move_button.isEnabled()
    monkeypatch.setattr(QInputDialog, "getItem", lambda *a: ("Retest", True))
    c.move_selected()
    assert [p.result.roll_number for p in w.store.pages(other)] == [files[0][2], files[2][2]]
    assert [s.result.roll_number for s in c.sheets] == [files[1][2]] and c.table.rowCount() == 1
    c.table.selectRow(0)
    monkeypatch.setattr(QMessageBox, "question", lambda *a: QMessageBox.Yes)
    c.delete_selected()
    assert c.sheets == [] and w.store.pages(c.exam.id) == [] and c.table.rowCount() == 0


def test_sections_and_a_short_key(app, tmp_path, monkeypatch):
    from omr_app import sections

    (path, answers, roll, set_code), = sheet_files(tmp_path, [4])
    w = window(tmp_path)
    c = w.check
    c.read_files([path])
    c.key_edit.setPlainText("".join(answers[q] or "A" for q in c.questions[:60]))  # only 60 questions in this paper
    assert "60 of 100 questions counted" in c.key_status.text()
    right = sum(1 for q in c.questions[:60] if answers[q])
    assert c.table.item(0, 6).text() == str(right) and c.full_marks() == 60

    def fill(dialog):
        dialog.add()
        dialog.table.item(0, 0).setText("Physics")
        dialog.table.item(0, 2).setText("30")
        dialog.table.item(0, 3).setText("4")
        dialog.table.item(0, 4).setText("1")  # minus 1 for wrong
        dialog.add()  # starts at 31
        return dialog.table.item(1, 1).text() == "31" and dialog.accept() is None
    monkeypatch.setattr(sections.SectionsDialog, "exec", fill)
    w.go("Marks and sections")
    c.sections_button.click()  # the real button, so a missing import or a broken dialog fails here
    assert [(s.name, s.wrong) for s in w.store.exam(c.exam.id).sections] == [("Physics", -1), ("Section 2", 0)]
    assert "Physics: questions 1 to 30" in c.sections_label.text()
    physics = sum(1 for q in c.questions[:30] if answers[q])
    assert c.table.item(0, 6).text() == str(4 * physics + right - physics)
    assert c.full_marks() == 4 * 30 + 30
    headers = [w.results.table.horizontalHeaderItem(i).text() for i in range(w.results.table.columnCount())]
    assert headers[4:7] == ["Score", "Physics", "Section 2"]
    assert w.results.table.item(0, 5).text() == str(4 * physics)

    dialog = sections.SectionsDialog(c, [], c.questions, (1, 0, 0))
    dialog.add()
    dialog.table.item(0, 2).setText("120")
    dialog.accept()
    assert "100" in dialog.error.text() and dialog.result() == 0  # stays open, says why


def test_clicking_an_exam_opens_its_page(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    from omr_core.sheet import SheetSpec, layout
    import json
    w = MainWindow(tmp_path / "omr.db")
    made = w.store.create_exam("Made here", json.dumps(layout(SheetSpec("Weekly", 30))))
    mapped = w.store.create_exam("Mapped", (TEMPLATES / "test-100.json").read_text(encoding="utf-8"))
    w.exams.refresh()
    assert "Check sheets" not in w.PAGES and w.showing() is w.exams
    row = next(r for r in range(w.exams.table.rowCount()) if w.exams.table.item(r, 0).text() == "Made here")
    w.exams.table.cellClicked.emit(row, 0)
    assert w.check.exam.id == made and w.showing() is w.print_sheet  # nothing read yet: print the sheet first
    assert w.exam_page.title.text() == "Made here" and w.print_sheet.save.isEnabled()
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(tmp_path / "sheet.pdf"), ""))
    w.print_sheet.save.click()
    assert (tmp_path / "sheet.pdf").read_bytes().startswith(b"%PDF")

    w.exam_settings.name.setText("Weekly 1")
    w.exam_settings.rename()
    assert w.store.exam(made).name == "Weekly 1" and w.exam_page.title.text() == "Weekly 1"
    w.exam_page.back.emit()
    assert w.showing() is w.exams

    w.open_exam(mapped)
    assert not w.print_sheet.save.isEnabled() and "print yourself" in w.print_sheet.note.text()
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)
    w.go("Exams")
    w.open_exam(mapped)
    w.go("Exam settings")
    assert w.showing() is w.exam_settings
    w.exam_settings.delete()
    assert w.check.exam is None and w.showing() is w.exams and [e[0] for e in w.store.exams()] == [made]


def test_the_installer_recipe_finds_everything_it_bundles():
    from pathlib import Path
    from omr_app import new_exam
    root = Path(__file__).parent.parent
    for path in ("src/omr_core/fonts/Poppins-Regular.ttf", "src/omr_app/icon.png", "templates/test-100.json",
                 "packaging/omr.ico", "LICENSE"):
        assert (root / path).is_file(), path
    assert (new_exam.TEMPLATES / "test-100.json").is_file()


def test_changing_the_sheet_reads_every_page_again(app, tmp_path, monkeypatch):
    import json
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    from omr_core.sheet import SheetSpec, layout
    files = sheet_files(tmp_path, [1, 2])
    t = load(TEMPLATES / "test-100.json")
    scanned = random_marks(t, 3)  # from the scanner: no file to go back to
    w = window(tmp_path)
    w.open_exam(w.store.create_exam("Wrong sheet", json.dumps(layout(SheetSpec("Weekly", 30)))))
    c = w.check
    c.read_files([f[0] for f in files])
    c.read_image(scan(draw(t, scanned[0]), seed=3), "Scanner", "Scanner", 1)
    assert all(s.result is None for s in c.sheets)

    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(TEMPLATES / "test-100.json"), ""))
    monkeypatch.setattr(QMessageBox, "question", lambda *a: QMessageBox.Yes)
    shown = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: shown.append(a[2]))
    w.exam_settings.change_sheet()
    assert c.exam.template.name == t.name and "3 of 3 pages read" in shown[0]
    assert [(s.result.roll_number, s.result.answers) for s in c.sheets] == [
        (files[0][2], files[0][1]), (files[1][2], files[1][1]), (scanned[2], scanned[1])]
    assert c.table.rowCount() == 3 and c.mismatch.isHidden()


def test_exams_are_backed_up_and_restored(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    shown = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: shown.append(a[2]))
    monkeypatch.setattr(QMessageBox, "warning", lambda *a: shown.append(a[2]))
    w = window(tmp_path)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a: (str(tmp_path / "backup.db"), ""))
    w.exams.backup()
    assert (tmp_path / "backup.db").is_file() and "Every exam is saved" in shown[-1]

    new_pc = MainWindow(tmp_path / "new" / "omr.db")
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a: (str(tmp_path / "backup.db"), ""))
    new_pc.exams.restore()
    assert shown[-1] == "Restored Mock." and new_pc.exams.table.item(0, 0).text() == "Mock"
    (tmp_path / "x.db").write_text("not a backup")
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a: (str(tmp_path / "x.db"), ""))
    new_pc.exams.restore()
    assert "isn't an OMR backup" in shown[-1] and new_pc.exams.table.rowCount() == 1


def test_a_printed_copy_is_checked_before_printing_the_rest(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    from omr_app.exam_page import print_check
    t = load(TEMPLATES / "test-100.json")
    assert "Every bubble reads empty" in print_check(scan(draw(t, {})), t)
    marks, answers, roll, _ = random_marks(t, 4)
    said = print_check(scan(draw(t, marks)), t)
    assert "This copy reads" in said and f"Roll number {roll}" in said
    assert "can't be read" in print_check(np.full((2339, 1654, 3), 255, np.uint8), t)

    w = window(tmp_path)
    path = tmp_path / "copy.png"
    cv2.imwrite(str(path), scan(draw(t, {})))
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a: (str(path), ""))
    shown = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: shown.append(a[2]))
    assert w.print_sheet.try_print.isEnabled()
    w.print_sheet.check_print()
    assert shown and shown[0].startswith("This copy reads")


def test_window_carries_the_app_name(app, tmp_path):
    w = window(tmp_path)
    assert w.windowTitle() == "OMR Checker  ·  Mock"
    w.open_exam(None)
    assert w.windowTitle() == "OMR Checker"
    # fits a 1920 x 1080 screen at 150% scaling (1280 x 720), with any page open
    for page in w.PAGES:
        if page not in w.LATER:
            w.go(page)
            assert w.minimumSizeHint().width() <= 1000 and w.minimumSizeHint().height() <= 700


def test_first_run_shows_the_steps_and_help_has_the_shortcuts(app, tmp_path, monkeypatch):
    from omr_app import exams, help
    w = MainWindow(tmp_path / "omr.db")
    assert not w.exams.start.isHidden() and w.exams.table.isHidden()
    assert "Create the exam" in w.exams.start.text()
    shown = []
    monkeypatch.setattr(help.HelpDialog, "exec", lambda self: shown.append(self.findChild(QLabel).text()))
    w.show_help()
    assert "Ctrl+N" in shown[0] and str(tmp_path) in shown[0] and "OMR Checker" in shown[0]
    monkeypatch.setattr(exams.NewExamWizard, "exec", lambda self: False)
    w.new_exam()
    assert w.showing() is w.exams
    w.import_scans()  # no exam open yet: stays on Exams
    assert w.showing() is w.exams

    w.open_exam(w.store.create_exam("Mock", (TEMPLATES / "test-100.json").read_text(encoding="utf-8")))
    w.exams.refresh()
    assert w.exams.start.isHidden() and not w.exams.table.isHidden()


def test_written_marks_are_typed_and_added_on_results(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QMessageBox, QTableWidgetItem
    files = sheet_files(tmp_path, [1, 2])
    w = window(tmp_path)
    c = w.check
    c.read_files([f[0] for f in files])
    c.key_edit.setPlainText("A" * 100)
    tab = w.written
    tab.refresh()
    assert tab.table.isHidden() and "Set the marks" in tab.note.text()
    tab.out_of.setValue(20)
    tab.save_out_of()
    tab.refresh()
    assert tab.table.rowCount() == 2 and "2 of 2 still to type" in tab.note.text()
    roll = tab.table.item(0, 0).text()
    before = next(c.score_of(s) for s in c.sheets if s.result.roll_number == roll)
    tab.table.setItem(0, 2, QTableWidgetItem("15"))
    assert w.store.written(c.exam.id) == {roll: 15}
    assert next(c.score_of(s) for s in c.sheets if s.result.roll_number == roll) == before + 15
    w.results.refresh()
    head = [w.results.table.horizontalHeaderItem(i).text() for i in range(w.results.table.columnCount())]
    assert head[4:7] == ["Score", "Sheet", "Written"]
    warned = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *a: warned.append(a[2]))
    tab.table.setItem(1, 2, QTableWidgetItem("25"))  # above the written part's 20
    assert warned and "0 to 20" in warned[0] and len(w.store.written(c.exam.id)) == 1


def test_answer_key_is_read_from_a_filled_master_sheet(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    t = load(TEMPLATES / "test-100.json")
    marks, answers, _, set_code = random_marks(t, 8)
    path = tmp_path / "master.png"
    cv2.imwrite(str(path), scan(draw(t, marks), seed=8))
    w = window(tmp_path)
    k = w.check.key_tab
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a: (str(path), ""))
    shown = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a: shown.append(a[2]))
    k.read_master()
    key = k.keys()[set_code]
    assert key == [answers[q] or "-" for q in w.check.questions]
    assert f"Key for set {set_code}" in shown[0] and k.set_code() == set_code


def test_absent_students_are_listed_and_saved(app, tmp_path, monkeypatch):
    from openpyxl import Workbook
    from PySide6.QtWidgets import QFileDialog, QMessageBox
    (path, _, roll, _), = sheet_files(tmp_path, [6])
    wb = Workbook()
    wb.active.append(["Roll number", "Name", "Class"])
    wb.active.append([roll, "Asha Rai", "10 A"])
    other = "9" * len(roll) if roll != "9" * len(roll) else "1" * len(roll)
    wb.active.append([other, "Bikash Shah", "10 A"])
    wb.save(tmp_path / "students.xlsx")
    w = window(tmp_path)
    w.check.read_files([path])
    monkeypatch.setattr(QFileDialog, "getOpenFileName", lambda *a, **k: (str(tmp_path / "students.xlsx"), ""))
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: None)
    w.students.import_button.click()
    s = w.students
    assert "2 students, 1 absent" in s.note.text() and s.save_absent.isEnabled()
    s.absent_only.setChecked(True)
    assert s.table.rowCount() == 1 and s.table.item(0, 1).text() == "Bikash Shah" and s.table.item(0, 3).text() == "Absent"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(tmp_path / "absent.csv"), ""))
    s.save_absent_list()
    assert (tmp_path / "absent.csv").read_text(encoding="utf-8-sig") == f"Roll number,Name,Class\n{other},Bikash Shah,10 A\n"


def test_numeric_questions_take_a_typed_key_and_a_typed_correction(app, tmp_path, monkeypatch):
    import json

    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QInputDialog

    from omr_core.sheet import SheetSection, SheetSpec, layout, render
    from sheets import fill
    d = layout(SheetSpec("JEE", 6, roll_digits=3, sets=0,
                         sections=[SheetSection("MCQ", 4), SheetSection("Numeric", 2, digits=3)]))
    t = parse(d)
    at = {g.kind if g.kind != "numeric" else g.first: i for i, g in enumerate(t.grids)}
    marks = {(at["roll_number"], 1, c): "pen" for c in range(3)}
    marks.update({(at["questions"], 0, 0): "pen",  # Q1 A
                  (at[5], 4, 0): "pen", (at[5], 2, 1): "pen",  # Q5 42
                  (at[6], 1, 0): "pen", (at[6], 7, 2): "pen"})  # Q6 1, a blank column, 7: for a person to read
    path = tmp_path / "jee.png"
    cv2.imwrite(str(path), scan(fill(render(d), t, marks)))
    w = MainWindow(tmp_path / "omr.db")
    w.open_exam(w.store.create_exam("JEE", json.dumps(d)))
    c = w.check
    c.read_files([str(path)])
    k, g = c.key_tab, c.key_tab.grid
    assert g.horizontalHeaderItem(2).text() == "A" and k.numeric == {4, 5}
    QTest.keyClick(g, Qt.Key_A)
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **kw: (kw["text"] + "2", True))
    g.setCurrentCell(4, 2)
    QTest.keyClick(g, Qt.Key_4)  # a digit starts the number
    assert c.key_edit.toPlainText() == "A, -, -, -, 42" and g.item(4, 2).text() == "42" and g.currentRow() == 5
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **kw: ("017", True))
    g.cellClicked.emit(5, 3)
    assert c.key_edit.toPlainText() == "A, -, -, -, 42, 17"
    assert c.score_of(c.sheets[0]) == 2  # Q6 was read as 1?7

    w.go("Review")
    assert [item for _, item, _ in w.review.items] == ["Q6"] and "gap" in w.review.why.text()
    edit, save = w.review.buttons.itemAt(0).widget(), w.review.buttons.itemAt(1).widget()
    assert edit.text() == "17"
    save.click()
    assert w.review.items == [] and c.score_of(c.sheets[0]) == 3


def test_new_exam_wizard_makes_numeric_sections(app, tmp_path):
    from omr_app.new_exam import NewExamWizard
    w = MainWindow(tmp_path / "omr.db")
    z = NewExamWizard(w, w.store)
    z.name.setText("JEE mock")
    z.add_section()
    z.add_section()
    z.sections.item(0, 1).setText("20")
    z.sections.item(1, 1).setText("5")
    z.sections.cellWidget(1, 3).setCurrentIndex(3)  # a number of up to 3 digits
    assert z.sections.item(1, 2).text() == "21 to 25" and z.total.value() == 25
    z.update_preview()
    assert [(g["first"], g["cols"]) for g in z.made["grids"] if g["kind"] == "numeric"] == [(q, 3) for q in range(21, 26)]
    z.shape.setCurrentIndex(z.shape.findData("circle"))
    z.labels.setCurrentIndex(z.labels.findData("numbers"))
    z.update_preview()
    assert z.made["print"]["bubble_shape"] == "circle" and z.made["bubble_mm"] == [2.4, 2.4]
    assert z.made["print"]["labels"] == "numbers"
    z.update_review()
    assert "answered with a number of up to 3 digits" in z.summary.text()


def test_a_sheet_printed_1_to_4_takes_its_key_by_those_numbers(app, tmp_path):
    import json

    from PySide6.QtTest import QTest

    from omr_core.sheet import SheetSpec, layout
    w = MainWindow(tmp_path / "omr.db")
    w.open_exam(w.store.create_exam("NEET", json.dumps(layout(SheetSpec("NEET", 10, sets=0, labels="numbers")))))
    k, g = w.check.key_tab, w.check.key_tab.grid
    assert [g.horizontalHeaderItem(c).text() for c in range(2, 6)] == ["A (1)", "B (2)", "C (3)", "D (4)"]
    QTest.keyClick(g, Qt.Key_2)
    QTest.keyClick(g, Qt.Key_D)
    assert w.check.key_edit.toPlainText() == "BD" and g.item(0, 3).text() == "B"


def test_designer_maps_numeric_answer_boxes(app, tmp_path):
    from omr_core.sheet import SheetSection, SheetSpec, layout, render
    d = layout(SheetSpec("JEE", 6, roll_digits=3, sets=0,
                         sections=[SheetSection("MCQ", 4), SheetSection("Numeric", 2, digits=3)]))
    t = parse(d)
    path = tmp_path / "blank.png"
    cv2.imwrite(str(path), scan(render(d)))
    w = window(tmp_path)
    w.go("Map printed sheet")
    m = w.designer
    m.load_scan(str(path))
    s = m.gray.shape[1] / t.page[0]
    for x, y in t.markers:
        m.mark_clicked(x * s, y * s)
    for g in t.grids:
        (x0, y0), (x1, y1) = g.centre(0, 0), g.centre(g.rows - 1, g.cols - 1)
        m.canvas.boxed.emit(QRectF((x0 - 4) * SCALE, (y0 - 4) * SCALE, (x1 - x0 + 8) * SCALE, (y1 - y0 + 8) * SCALE))
    assert [(b["kind"], b.get("first")) for b in m.blocks] == [(g.kind, g.first if g.kind != "roll_number" else None)
                                                              for g in t.grids]
    assert parse(m.template_dict()).questions() == [1, 2, 3, 4, 5, 6]


def test_ready_made_layouts_fill_the_wizard_and_fit_the_sheet(app, tmp_path):
    from omr_app import new_exam
    from omr_core.store import Store
    z = new_exam.NewExamWizard(None, Store(tmp_path / "omr.db"))
    for name, (sections, total, options, right, wrong, roll, sets, labels, shape, _) in new_exam.READY.items():
        z.ready.setCurrentText(name)
        assert z.total.value() == (sum(s[1] for s in sections) or total) and z.options.value() == options
        assert (z.roll.value(), z.labels.currentData(), z.shape.currentData()) == (roll, labels, shape)
        assert [(s.name, s.digits) for s in z.read_sections()] == [(s[0], (s[2:] or [0])[0]) for s in sections]
        assert [(s.correct, s.wrong) for s in z.sheet_sections()] == [(right, wrong)] * len(sections)
        z.update_preview()
        assert z.made is not None, (name, z.layout_problem)  # fits one page
        assert len([g for g in z.made["grids"] if g["kind"] == "set_code"]) == bool(sets)
    z.ready.setCurrentText(next(n for n in new_exam.READY if "JEE" in n))
    z.update_preview()
    assert [g["first"] for g in z.made["grids"] if g["kind"] == "numeric"] == [*range(21, 26), *range(46, 51),
                                                                               *range(71, 76)]


def test_a_test_series_shares_students_and_ranks_them_together(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QInputDialog, QMessageBox

    from omr_app import series
    from omr_app.new_exam import NewExamWizard
    from omr_core.reader import PageResult
    from omr_core.students import Student
    w = window(tmp_path)  # "Mock", in no series yet
    store, first = w.store, w.check.exam.id
    store.set_students(first, [Student("11111", "Asha", "12")])
    assert w.exams.series.isHidden()
    w.exam_settings.series.setCurrentText("JEE 2027")
    w.exam_settings.set_series()
    assert store.exam(first).details["series"] == "JEE 2027" and "test series JEE 2027" in w.exam_page.about.text()
    z = NewExamWizard(w, store)
    z.name.setText("Mock 2")
    assert z.series.findText("JEE 2027") >= 0
    z.series.setCurrentText("JEE 2027")
    z.update_review()
    z.create()
    assert store.students(z.exam_id) == store.students(first)  # carried over, not imported again

    for exam, right in ((first, 30), (z.exam_id, 40)):
        store.save_scheme(exam, "A" * 100 if exam == first else "", 1, 0, 0)
        store.add_page(exam, "b", "s.png", 1, PageResult("11111", "A", {q: "A" if q <= right else "B"
                                                                        for q in range(1, 101)}, []))
    w.exams.refresh()
    assert not w.exams.series.isHidden()
    d = series.SeriesDialog(w, store)
    assert d.series.currentText() == "JEE 2027" and d.table.rowCount() == 1
    assert [d.table.item(0, c).text() for c in range(d.table.columnCount())] == [
        "1", "11111", "Asha", "12", "30", "absent", "30", "1", "30"]
    assert "Not counted: Mock 2: no results yet" in d.note.text()
    out = tmp_path / "series.xlsx"
    monkeypatch.setattr(series.QFileDialog, "getSaveFileName", lambda *a: (str(out), ""))
    monkeypatch.setattr(series.QMessageBox, "information", lambda *a: None)
    d.save_excel()
    assert out.exists()

    other = store.create_exam("Quiz", (TEMPLATES / "test-100.json").read_text(encoding="utf-8"))
    w.open_exam(other)
    monkeypatch.setattr(QInputDialog, "getItem", lambda _p, _t, _l, items, *a: (items[0], True))
    monkeypatch.setattr(QMessageBox, "question", lambda *a: QMessageBox.Yes)
    w.students.copy_students()
    assert store.students(other) == store.students(first) and w.students.table.rowCount() == 1


def test_topics_are_typed_once_and_reach_the_results(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QInputDialog, QMessageBox
    w = window(tmp_path)
    c = w.check
    answers = iter([("1-50: Algebra\nGeometry", True), ("1-50: Algebra\n51-100: Geometry", True)])
    monkeypatch.setattr(QInputDialog, "getMultiLineText", lambda *a: next(answers))
    warned = []
    monkeypatch.setattr(QMessageBox, "warning", lambda *a: warned.append(a[2]))
    c.edit_topics()  # a line without questions is refused, and the text comes back to fix
    assert "line 2" in warned[0] and c.topics[51] == "Geometry"
    assert w.store.topics(c.exam.id) == c.topics and c.topics_label.text().startswith("2 topics over 100 questions")
    path, answers_given, _, _ = sheet_files(tmp_path, [3])[0]
    c.read_files([path])
    c.key_edit.setPlainText("".join(answers_given[q] or "A" for q in c.questions))
    w.go("Results")
    g = w.results.graded[0]
    assert set(g.topics) == {"Algebra", "Geometry"} and sum(n for _, n in g.topics.values()) == 100


def test_templates_page_saves_edits_as_a_new_template_and_never_changes_a_built_in(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QInputDialog, QMessageBox

    from omr_app import new_exam
    w = MainWindow(tmp_path / "omr.db")
    w.go("Templates")
    p = w.templates
    neet = next(n for n in new_exam.READY if "NEET" in n)
    p.list.setCurrentItem(p.list.findItems(neet, Qt.MatchExactly)[0])
    assert p.current == neet and not p.save_button.isVisibleTo(p) and not p.delete_button.isVisibleTo(p)
    p.sections.item(0, 1).setText("50")  # Physics 45 -> 50
    p.shape.setCurrentIndex(p.shape.findData("oval"))
    p.update_preview()
    assert p.edited() and not p.preview.pixmap().isNull()
    monkeypatch.setattr(QMessageBox, "question", lambda *a: QMessageBox.No)
    p.list.setCurrentRow(0)  # keep the changes: stays on NEET
    assert p.current == neet and p.list.currentItem().text() == neet and p.edited()
    p.reset()  # Reset to default
    assert not p.edited() and p.sections.item(0, 1).text() == "45"

    p.sections.item(0, 1).setText("50")
    p.wrong.setValue(-0.5)
    p.track.setChecked(True)
    p.update_preview()
    assert p.bar.numbering.isVisibleTo(p.bar) and p.bar.save.isEnabled()  # numbered copies print from here too
    names = iter([(neet, True), ("NEET mock, Physics 50", True)])  # a built-in's name is refused, then a new one
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: next(names))
    monkeypatch.setattr(QMessageBox, "warning", lambda *a: None)
    p.save_as()
    assert p.current == "NEET mock, Physics 50" and p.save_button.isVisibleTo(p)
    assert new_exam.READY[neet].sections[0] == ("Physics", 45)  # the built-in is untouched
    mine = new_exam.saved_templates(w.store)[0]["NEET mock, Physics 50"]
    assert mine.sections[0] == ("Physics", 50) and mine.wrong == -0.5 and mine.labels == "numbers" and mine.track

    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("Weekly NEET", True))
    p.rename()
    monkeypatch.setattr(QInputDialog, "getText", lambda *a, **k: ("Weekly NEET 2", True))
    p.duplicate()
    assert list(w.store.templates()) == ["Weekly NEET", "Weekly NEET 2"] and p.current == "Weekly NEET 2"

    z = new_exam.NewExamWizard(w, w.store)  # listed under Start from, and fills the wizard
    z.ready.setCurrentText("Weekly NEET")
    assert z.read_sections()[0].questions == 50 and z.total.value() == 185 and z.wrong.value() == -0.5

    monkeypatch.setattr(QMessageBox, "question", lambda *a: QMessageBox.Yes)
    p.delete()  # the copy, open after duplicating
    assert list(w.store.templates()) == ["Weekly NEET"] and p.current == next(iter(new_exam.READY))
    w.store.save_template("Broken", "{not json")
    p.refresh("Broken")
    assert "Broken" in p.damaged.text() and "Broken" not in p.mine


def test_a_saved_template_is_checked_before_use():
    import json

    from omr_app.new_exam import READY, Ready, ready_from_json
    for r in READY.values():
        assert ready_from_json(json.dumps(r._asdict())) == r._replace(sections=[tuple(s) for s in r.sections])
    good = Ready([("Physics", 45)], 0, 4, 4, -1)._asdict()
    for bad in ({"options": 9}, {"roll": 20}, {"sets": 1}, {"labels": "roman"}, {"shape": "star"},
                {"sections": [["", 5]]}, {"sections": [["A", 0]]}, {"sections": []}, {"total": "100"},
                {"track": "yes"}):
        with pytest.raises(ValueError):
            ready_from_json(json.dumps({**good, **bad}))
    for text in ("[]", "{}", "nonsense", json.dumps({**good, "extra": 1})):
        with pytest.raises(ValueError):
            ready_from_json(text)


def test_numbered_sheets_print_and_their_numbers_show_where_sheets_are_listed(app, tmp_path, monkeypatch):
    import csv
    import json

    import pypdfium2 as pdfium
    from PySide6.QtPrintSupport import QPrinter
    from PySide6.QtWidgets import QFileDialog

    from omr_core.sheet import SheetSpec, layout
    w = MainWindow(tmp_path / "omr.db")
    w.open_exam(w.store.create_exam("Weekly", json.dumps(layout(SheetSpec("Weekly", 30, roll_digits=4, sets=0,
                                                                           track=True)))))
    bar = w.print_sheet.bar
    assert bar.numbering.isVisibleTo(bar)
    bar.prefix.setText("W-")
    bar.first.setValue(7)
    bar.copies.setValue(4)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(tmp_path / "sheets.pdf"), ""))
    bar.save.click()
    pdf = pdfium.PdfDocument(str(tmp_path / "sheets.pdf"))
    pages = [pdf[k].render(scale=200 / 72).to_numpy().copy() for k in range(len(pdf))]
    pdf.close()
    assert len(pages) == 4
    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(str(tmp_path / "printed.pdf"))
    bar.paint(printer, bar.numbers()[:2])  # what Print… sends once the printer is chosen
    printed = pdfium.PdfDocument(str(tmp_path / "printed.pdf"))
    assert len(printed) == 2
    printed.close()

    files = []
    for k in (0, 1, 1, 3):  # W-0008 scanned twice, W-0009 never
        files.append(str(tmp_path / f"scan{len(files)}.png"))
        cv2.imwrite(files[-1], scan(pages[k], angle=0.5, seed=k))
    c = w.check
    c.read_files(files)
    column = c.table.columnCount() - 1
    assert not c.table.isColumnHidden(column)
    assert [c.table.item(r, column).text() for r in range(4)] == ["W-0007", "W-0008", "W-0008", "W-0010"]
    assert c.sheet_numbers.text() == "Sheet number W-0008 on more than one page. Sheet number W-0009 not scanned."
    c.key_edit.setPlainText("A" * 30)
    results = w.results.table
    assert results.horizontalHeaderItem(results.columnCount() - 1).text() == "Sheet no."
    assert {results.item(r, results.columnCount() - 1).text() for r in range(4)} == {"W-0007", "W-0008", "W-0010"}
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(tmp_path / "results.csv"), ""))
    monkeypatch.setattr(ColumnsDialog, "exec", lambda self: 1)  # every column, in the default order
    c.save_csv()
    rows = list(csv.reader(open(tmp_path / "results.csv", encoding="utf-8-sig")))
    at = rows[0].index("Sheet no.")
    assert [r[at] for r in rows[1:]] == ["W-0007", "W-0008", "W-0008", "W-0010"]
    from openpyxl import load_workbook
    from PySide6.QtWidgets import QMessageBox
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.Yes)  # blank roll numbers to check
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(tmp_path / "merit.xlsx"), ""))
    w.results.excel.click()
    merit = list(load_workbook(tmp_path / "merit.xlsx")["Merit list"].values)
    assert merit[1][-1] == "Sheet no." and {r[-1] for r in merit[2:]} == {"W-0007", "W-0008", "W-0010"}


def test_csv_columns_are_picked_ordered_and_remembered_per_exam(app, tmp_path, monkeypatch):
    import csv

    from PySide6.QtWidgets import QFileDialog

    from omr_app.columns import arrange

    assert arrange(["a", "b", "c"], [["c", False], ["gone", True], ["a", True], ["c", True]]) == [
        ("c", True), ("a", True), ("b", True)]
    w = window(tmp_path)
    c = w.check
    page = np.full((900, 640), 255, np.uint8)
    path = tmp_path / "blank.png"
    cv2.imwrite(str(path), page)
    c.read_files([str(path)])
    drain(w.scan)
    seen = []

    def choose(self):
        seen.append(self.choice())
        names = [n for n, _ in self.choice()]
        if len(seen) == 1:  # File first, then Roll number; Review unticked
            self.fill([("File", True), ("Roll number", True)] + [(n, n != "Review") for n in names
                                                                  if n not in ("File", "Roll number")])
            self.tick_all(False)
            assert not self.save.isEnabled()  # nothing to save
            self.fill([("File", True), ("Roll number", True), ("Review", False)])
        return 1

    monkeypatch.setattr(ColumnsDialog, "exec", choose)
    out = tmp_path / "picked.csv"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(out), ""))
    w.results.csv.click()
    rows = list(csv.reader(open(out, encoding="utf-8-sig")))
    assert rows[0] == ["File", "Roll number"] and rows[1][0] == str(path)
    c.save_csv()  # the same exam starts from the last choice
    assert seen[1][:3] == [("File", True), ("Roll number", True), ("Review", False)]
    assert ("Rank", True) in seen[1]  # columns left out of the last choice come back ticked at the end


def test_cut_off_corner_shows_why_live_and_on_review(app, tmp_path):
    import json

    from omr_core.sheet import SheetSpec, layout, render
    from omr_core.template import parse as parse_template
    from sheets import fill

    spec = layout(SheetSpec("Mock test", 200, options=4, roll_digits=6, sets=4))
    t = parse_template(spec)
    path = tmp_path / "cut.png"
    cv2.imwrite(str(path), scan(fill(render(spec), t, random_marks(t, 1)[0]), 6, seed=1))  # a corner mark half lost
    w = MainWindow(tmp_path / "omr.db")
    w.open_exam(w.store.create_exam("Cut", json.dumps(spec)))
    w.scan.read_files([str(path)])
    drain(w.scan)
    assert w.scan.latest.text().startswith("Can't read this page: a corner mark is cut off at the edge of the scan.")
    assert w.review.items and w.review.why.text() == "unreadable: a corner mark is cut off at the edge of the scan"


def test_imported_sheets_show_live_with_their_marks_and_can_pause_and_stop(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QFileDialog

    from omr_app.scan import RIGHT, WRONG, marked_sheet
    from omr_core.reader import SCALE
    files = sheet_files(tmp_path, [7, 8, 9, 7])  # the last is sheet 7 again: the same roll number twice
    w = window(tmp_path)
    s, c = w.scan, w.check
    path, answers, roll, set_code = files[0]
    c.key_edit.setPlainText(f"{set_code}: " + "".join(answers[q] or "A" for q in c.questions))
    monkeypatch.setattr(QFileDialog, "getOpenFileNames", lambda *a, **k: ([f[0] for f in files], ""))
    w.import_scans()  # Ctrl+O reads on the Scan tab
    assert w.showing() is s and s.reading_files and s.pause_button.isEnabled()
    s.step()
    assert s.live.rowCount() == 1 and s.live.item(0, 1).text() == roll and not s.latest.pixmap().isNull()
    assert roll in s.latest_text.text() and "Score" in s.latest_text.text()
    s.toggle_pause()
    s.step()  # paused: nothing more is read
    assert s.live.rowCount() == 1 and s.pause_button.text() == "Resume"
    s.toggle_pause()
    drain(s)
    assert s.live.rowCount() == 4 and s.read.value.text() == "4" and s.same_roll.value.text() == "2"
    assert s.status.text().startswith("Done.") and not s.running and s.speed.value.text() != "–"
    assert [s.live.item(r, 1).text() for r in range(4)] == [f[2] for f in files]
    s.live.cellClicked.emit(1, 0)  # an earlier sheet comes back on the live view
    assert files[1][2] in s.latest_text.text()

    # The marks drawn: a ring in green on a right answer, in red on a wrong one.
    t = c.exam.template
    stored = c.sheets[0]
    q = next(q for q in c.questions if answers[q])
    g = t.grid_for(q)
    x, y = g.centre(q - g.first, g.options.index(answers[q]))
    radius = round(max(t.bubble) / 2 * SCALE) + 3
    on_ring = (round(y * SCALE), round(x * SCALE) + radius)
    image = w.store.image(stored.id)
    assert tuple(marked_sheet(image, t, stored.result, {q: answers[q]})[on_ring]) == RIGHT
    other = next(o for o in g.options if o != answers[q])
    assert tuple(marked_sheet(image, t, stored.result, {q: other})[on_ring]) == WRONG

    # Stopping part-way keeps what was read and says what wasn't.
    s.read_files([f[0] for f in files[:3]])
    s.step()
    s.stop()
    assert s.live.rowCount() == 1 and "2 files not read" in s.status.text() and not s.running


def test_start_scanning_is_one_click_from_the_dashboard_and_the_exam_page(app, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QMessageBox
    t = load(TEMPLATES / "test-100.json")
    marks, _, roll, _ = random_marks(t, 4)
    monkeypatch.setattr(scanner, "scanners", lambda: [FakeScanner([scan(draw(t, marks))])])
    w = window(tmp_path)
    w.scan.find_scanners()
    assert "Scanner ready: Test scanner" in w.scan.ready.text() and "Start scanning" in w.scan.start_button.text()
    w.go("Dashboard")
    w.dashboard.scan_wanted.emit()
    assert w.showing() is w.scan and w.scan.thread is not None
    w.scan.thread.wait()
    for _ in range(5):
        app.processEvents()
    assert [p.result.roll_number for p in w.store.pages(w.check.exam.id)] == [roll]

    # No scanner and no folder: it says what to do and offers the other ways in.
    monkeypatch.setattr(scanner, "scanners", lambda: [])
    w.scan.find_scanners()
    assert w.scan.ready.text() == "No scanner found through Windows"
    w.scan.folder.setText("")
    asked, imported = [], []
    monkeypatch.setattr(QMessageBox, "exec", lambda box: asked.append(box.text()))
    monkeypatch.setattr(QMessageBox, "clickedButton",
                        lambda box: next(b for b in box.buttons() if b.text() == "Import scans…"))
    monkeypatch.setattr(w.scan, "import_files", lambda: imported.append(True))
    w.go("Results")
    w.exam_page.scan_wanted.emit()
    assert w.showing() is w.scan and "No scanner was found" in asked[0] and imported == [True]


def test_portable_copy_keeps_its_data_beside_it(tmp_path, monkeypatch):
    import sys

    from omr_app.paths import data_dir, portable_dir

    assert portable_dir() is None  # run from source
    monkeypatch.setattr(sys, "frozen", True, raising=False)
    monkeypatch.setattr(sys, "executable", str(tmp_path / "OMR.exe"))
    assert portable_dir() is None and data_dir().name == "OMR"  # installed: %LOCALAPPDATA%\OMR
    (tmp_path / "portable.txt").write_text("")
    assert portable_dir() == data_dir() == tmp_path / "data"
